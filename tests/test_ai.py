"""Test module AI Search & GEO (R1), theo spec §8.

FakeGA4 giả lập GA4 Data API ở mức từng phiên: tự áp dimensionFilter và tự tính metric như GA4 (rate từ số phiên,
không cộng rate). Nhờ vậy test kiểm được cả bộ lọc tool gửi đi lẫn cách tool đọc số.
"""
import json
import re
import unittest
from datetime import date

from googleapiclient.errors import HttpError

from seo_report import ai_classifier as C
from seo_report import db
from seo_report.ai_analytics import AIAnalytics, Filters, change, previous_period, _bucket
from seo_report.ga4_query import GA4Client

D0 = date(2026, 9, 1)


def sess(source, medium, channel, n=1, dur=60, engaged=True, events=None, country="Vietnam", page="/", day=D0):
    return [{"sessionSource": source, "sessionMedium": medium, "sessionDefaultChannelGroup": channel,
             "country": country, "landingPage": page, "date": day.strftime("%Y%m%d"), "dur": dur,
             "engaged": engaged, "events": dict(events or {})} for _ in range(n)]


def match(expr, item) -> bool:
    if "orGroup" in expr:
        return any(match(e, item) for e in expr["orGroup"]["expressions"])
    if "andGroup" in expr:
        return all(match(e, item) for e in expr["andGroup"]["expressions"])
    if "notExpression" in expr:
        return not match(expr["notExpression"], item)
    f = expr["filter"]
    value = str(item.get(f["fieldName"], ""))
    if "inListFilter" in f:
        vals = f["inListFilter"]["values"]
        if f["inListFilter"].get("caseSensitive", True):
            return value in vals
        return value.lower() in [v.lower() for v in vals]
    sf = f["stringFilter"]
    flags = 0 if sf.get("caseSensitive", False) else re.I
    if sf["matchType"] == "EXACT":
        return value.lower() == sf["value"].lower() if flags else value == sf["value"]
    if sf["matchType"] == "BEGINS_WITH":
        return value.lower().startswith(sf["value"].lower())
    if sf["matchType"] == "FULL_REGEXP":
        return re.fullmatch(sf["value"], value, flags) is not None
    if sf["matchType"] == "PARTIAL_REGEXP":
        return re.search(sf["value"], value, flags) is not None
    raise ValueError(sf["matchType"])


class Req:
    def __init__(self, fn):
        self.fn = fn

    def execute(self):
        return self.fn()


class FakeGA4:
    def __init__(self, sessions):
        self.sessions, self.calls = sessions, []

    def properties(self):
        return self

    def runReport(self, property, body):
        return Req(lambda: self._run(body))

    def _run(self, body):
        self.calls.append(body)
        dims = [d["name"] for d in body["dimensions"]]
        mets = [m["name"] for m in body["metrics"]]
        for m in mets:
            if m.startswith("sessionKeyEventRate:") and " " in m:
                resp = type("R", (), {"status": 400, "reason": "Bad"})()
                raise HttpError(resp, b'{"error":{"message":"Field is not a valid metric"}}')
        r0 = body["dateRanges"][0]
        lo, hi = r0["startDate"].replace("-", ""), r0["endDate"].replace("-", "")
        items = [s for s in self.sessions if lo <= s["date"] <= hi]
        if "eventName" in dims:  # tách mỗi event một dòng
            items = [{**s, "eventName": ev, "_count": n} for s in items for ev, n in s["events"].items()]
        if body.get("dimensionFilter"):
            items = [s for s in items if match(body["dimensionFilter"], s)]
        groups = {}
        for s in items:
            groups.setdefault(tuple(s.get(d) for d in dims), []).append(s)
        rows = []
        for key, ss in groups.items():
            n = len(ss)
            values = []
            for m in mets:
                if m == "sessions":
                    v = n
                elif m == "engagedSessions":
                    v = sum(s["engaged"] for s in ss)
                elif m == "engagementRate":
                    v = sum(s["engaged"] for s in ss) / n
                elif m == "userEngagementDuration":
                    v = sum(s["dur"] for s in ss)
                elif m == "averageSessionDuration":
                    v = sum(s["dur"] for s in ss) / n
                elif m == "keyEvents":
                    v = sum(s.get("_count", sum(s["events"].values())) for s in ss)
                elif m == "sessionKeyEventRate":
                    v = sum(1 for s in ss if sum(s["events"].values())) / n
                elif m.startswith("sessionKeyEventRate:"):
                    v = sum(1 for s in ss if s["events"].get(m.split(":", 1)[1])) / n
                else:
                    raise ValueError(m)
                values.append({"value": str(v)})
            rows.append({"dimensionValues": [{"value": k} for k in key], "metricValues": values})
        # GA4 thật bỏ dòng có mọi metric bằng 0.
        rows = [r for r in rows if any(float(v["value"]) for v in r["metricValues"])]
        if body.get("orderBys"):
            idx = mets.index(body["orderBys"][0]["metric"]["metricName"])
            rows.sort(key=lambda r: -float(r["metricValues"][idx]["value"]))
        off, lim = body.get("offset", 0), body.get("limit", 10000)
        return {"dimensionHeaders": [{"name": d} for d in dims], "metricHeaders": [{"name": m} for m in mets],
                "rows": rows[off:off + lim], "rowCount": len(rows), "metadata": {}}


def spec_dataset():
    """Dữ liệu theo ví dụ trong spec §7.2/§8: 609 phiên AI, 17 key events, thời lượng theo nguồn 229/190/668/50 giây."""
    s = []
    s += sess("chatgpt.com", "ai-assistant", "AI Assistant", 537, dur=229)
    s += sess("gemini.google.com", "ai-assistant", "AI Assistant", 66, dur=190)
    s += sess("claude.ai", "ai-assistant", "AI Assistant", 3, dur=668, page="/blog-a/")
    s += sess("perplexity.ai", "ai-assistant", "AI Assistant", 3, dur=50, engaged=False)
    for i in range(15):  # 15 phiên có key event, 2 phiên trong đó có 2 event → 17 key events
        s[i]["events"] = {"booking_success": 2} if i < 2 else {"Click Zalo": 1}
    s += sess("chatgpt", "paid", "Paid Other", 113, events={"purchase": 1})          # quảng cáo: loại
    s += sess("google", "organic", "Organic Search", 6000, dur=100)
    s += sess("facebook.com", "referral", "Organic Social", 400, dur=40)
    s += sess("(direct)", "(none)", "Direct", 2000)
    return s


def make(sessions, mode=C.HYBRID):
    return AIAnalytics(GA4Client(FakeGA4(sessions), "properties/1"), mode=mode, today=date(2026, 12, 1))


class ClassifierTest(unittest.TestCase):
    def test_native_singular_and_plural(self):
        for ch in ("AI Assistant", "AI Assistants", "ai assistant"):
            self.assertEqual(C.classify("x.com", "referral", ch)[::2], (C.AI_ASSISTANTS, C.NATIVE))

    def test_hosts_aliases_and_boundaries(self):
        self.assertEqual(C.classify("chat.chatgpt.com", "referral", "Referral")[:2], (C.AI_ASSISTANTS, "CHATGPT"))
        self.assertEqual(C.classify("WWW.Perplexity.AI", "referral", "Referral")[1], "PERPLEXITY")
        self.assertEqual(C.classify("copilot.com", "referral", "Referral")[1], "COPILOT")
        self.assertEqual(C.classify("gemini", "(not set)", "Unassigned")[:2], (C.AI_ASSISTANTS, "GEMINI"))
        self.assertEqual(C.classify("evilchatgpt.com", "referral", "Referral")[0], C.OTHER)
        self.assertEqual(C.classify("geminiapp", "referral", "Referral")[0], C.OTHER)

    def test_google_never_ai(self):
        for src in ("google", "google.com", "www.google.com"):
            self.assertEqual(C.classify(src, "organic", "Organic Search")[0], C.ORGANIC_SEARCH)

    def test_paid_excluded(self):
        for medium in ("paid", "cpc", "ppc", "paid_social", "retargeting"):
            self.assertEqual(C.classify("chatgpt", medium, "Paid Other")[0], C.OTHER)

    def test_modes(self):
        self.assertEqual(C.classify("gemini", "(not set)", "Unassigned", C.NATIVE_ONLY)[0], C.OTHER)
        self.assertEqual(C.classify("x.com", "referral", "AI Assistant", C.FALLBACK_ONLY)[0], C.OTHER)

    def test_ga4_filter_equals_python_classifier(self):
        """Bộ lọc gửi GA4 và classify() phải chọn đúng cùng tập phiên, ở mọi chế độ."""
        data = spec_dataset() + sess("gemini", "(not set)", "Unassigned", 11) + \
            sess("chatgpt.com", "zalo", "Unassigned", 2) + sess("newai.example", "ai-assistant", "AI Assistant", 4) + \
            sess("evilchatgpt.com", "referral", "Referral", 5)
        for mode in (C.HYBRID, C.NATIVE_ONLY, C.FALLBACK_ONLY):
            f = C.ga4_filter(mode)
            by_filter = sum(1 for s in data if match(f, s))
            by_python = sum(1 for s in data if C.classify(s["sessionSource"], s["sessionMedium"],
                                                            s["sessionDefaultChannelGroup"], mode)[0] == C.AI_ASSISTANTS)
            self.assertEqual(by_filter, by_python, mode)
        self.assertEqual(sum(1 for s in data if match(C.ga4_filter(C.HYBRID), s)), 609 + 11 + 2 + 4)
        self.assertEqual(sum(1 for s in data if match(C.paid_ai_filter(), s)), 113)


class AnalyticsTest(unittest.TestCase):
    def setUp(self):
        self.ai = make(spec_dataset())
        self.end = D0

    def test_summary_uses_aggregate_not_sums(self):
        env = self.ai.summary(D0, D0, compare=False)
        a = env["data"]["ai"]
        self.assertEqual((a["sessions"], a["keyEvents"]), (609, 17))
        self.assertAlmostEqual(a["avgEngagementTimePerSession"], 137667 / 609)  # ≈ 226 giây, không phải 1137
        self.assertAlmostEqual(a["keyEventsPer100Sessions"], 100 * 17 / 609)    # 2,79
        self.assertAlmostEqual(a["sessionKeyEventRate"], 15 / 609)              # 2,46%: KHÁC 17/609
        self.assertAlmostEqual(a["engagementRate"], 606 / 609)
        self.assertEqual(env["data"]["paidAi"]["sessions"], 113)
        share = env["data"]["share"]
        self.assertEqual(share["allSessions"], 609 + 113 + 6000 + 400 + 2000)
        self.assertEqual(share["organicBase"], 609 + 6000 + 400)
        # KPI tổng đến từ truy vấn không có dimension.
        body = self.ai.client.data.calls[0]
        self.assertEqual(body["dimensions"], [])

    def test_sources_rates_per_source_from_own_query(self):
        env = self.ai.sources(D0, D0, compare=False)
        rows = {r["source"]: r for r in env["data"]["sources"]}
        self.assertEqual([r["source"] for r in env["data"]["sources"]], ["CHATGPT", "GEMINI", "CLAUDE", "PERPLEXITY"])
        self.assertAlmostEqual(rows["CHATGPT"]["avgEngagementTimePerSession"], 229)
        self.assertAlmostEqual(rows["PERPLEXITY"]["engagementRate"], 0.0)
        self.assertEqual(env["data"]["totalSessions"], 609)
        self.assertNotIn("GA4_TOTAL_DISCREPANCY", [w["code"] for w in env["warnings"]])

    def test_unknown_native_source_flagged(self):
        env = make(spec_dataset() + sess("newai.example", "ai-assistant", "AI Assistant", 4)).sources(D0, D0, compare=False)
        other = [r for r in env["data"]["sources"] if r["source"] == "OTHER_AI"][0]
        self.assertEqual((other["sessions"], other["engagementRate"]), (4, None))
        self.assertIn("UNMAPPED_AI_SOURCE", [w["code"] for w in env["warnings"]])

    def test_zero_baseline_and_low_sample(self):
        data = sess("chatgpt.com", "ai-assistant", "AI Assistant", 5, day=date(2026, 9, 2))
        env = make(data).summary(date(2026, 9, 2), date(2026, 9, 2))
        self.assertIsNone(env["data"]["aiDeltas"]["sessions"])
        codes = [w["code"] for w in env["warnings"]]
        self.assertIn("NO_BASELINE", codes)
        self.assertIn("LOW_SAMPLE_SIZE", codes)
        self.assertEqual(env["status"], "PARTIAL")

    def test_no_sessions_means_no_rate(self):
        env = make(sess("google", "organic", "Organic Search", 10)).summary(D0, D0, compare=False)
        self.assertEqual(env["data"]["ai"]["sessions"], 0)
        self.assertIsNone(env["data"]["ai"]["engagementRate"])
        self.assertEqual(env["data"]["share"]["ofAllSessions"], 0.0)  # có mẫu số (10 phiên) nên 0% là số thật

    def test_filters_propagate_to_every_query(self):
        self.ai.summary(D0, D0, Filters(country="Vietnam"), compare=False)
        for body in self.ai.client.data.calls:
            self.assertIn('"country"', json.dumps(body.get("dimensionFilter", {})))

    def test_landing_pages_keep_not_set(self):
        data = spec_dataset() + sess("chatgpt.com", "ai-assistant", "AI Assistant", 2, page="(not set)")
        env = make(data).landing_pages(D0, D0, compare=False)
        rows = {r["landingPage"]: r for r in env["data"]["rows"]}
        self.assertTrue(rows["(not set)"]["notSet"])
        self.assertEqual(rows["/blog-a/"]["sessions"], 3)
        self.assertIn("DATA_NOT_AVAILABLE", [w["code"] for w in env["warnings"]])

    def test_key_events_breakdown_and_incompatible_metric(self):
        env = self.ai.key_events(D0, D0, ["booking_success", "Click Zalo", "purchase"])
        ev = {e["event"]: e for e in env["data"]["events"]}
        self.assertEqual((ev["booking_success"]["keyEvents"], ev["Click Zalo"]["keyEvents"]), (4, 13))
        self.assertEqual(ev["purchase"]["keyEvents"], 0)  # purchase chỉ có ở traffic quảng cáo, đã loại
        self.assertEqual(ev["purchase"]["sessionKeyEventRate"], 0.0)  # 0% thật, không phải "không có số"
        self.assertAlmostEqual(ev["booking_success"]["sessionKeyEventRate"], 2 / 609)
        self.assertIsNone(ev["Click Zalo"]["sessionKeyEventRate"])
        self.assertIn("METRIC_INCOMPATIBLE", [w["code"] for w in env["warnings"]])

    def test_permission_error_is_error_status_not_zero(self):
        class Denied(FakeGA4):
            def _run(self, body):
                raise HttpError(type("R", (), {"status": 403, "reason": "Forbidden"})(), b"{}")
        env = AIAnalytics(GA4Client(Denied([]), "properties/1"), today=date(2026, 12, 1)).summary(D0, D0)
        self.assertEqual((env["status"], env["data"]), ("ERROR", {}))
        self.assertEqual(env["warnings"][0]["code"], "API_PERMISSION_DENIED")


class R3Test(unittest.TestCase):
    def setUp(self):
        data = spec_dataset()
        for s in data:
            if s["sessionSource"] == "gemini.google.com":
                s["landingPage"], s["country"] = "/blog-b/", "Japan"
        data += sess("google", "organic", "Organic Search", 50, page="/blog-a/")
        self.ai = make(data)

    def test_segments_rates_from_aggregate(self):
        env = self.ai.segments(D0, D0, {"Blog": ["/blog-a/", "/blog-b/"], "Trang chủ": ["/"]}, compare=False)
        seg = {x["segment"]: x for x in env["data"]["segments"]}
        self.assertEqual((seg["Blog"]["sessions"], seg["Trang chủ"]["sessions"]), (69, 540))
        # engagement của nhóm Blog = 69/69 phiên engaged, tính trên đúng nhóm (không lấy trung bình các trang)
        self.assertAlmostEqual(seg["Blog"]["engagementRate"], 69 / 69)
        self.assertAlmostEqual(seg["Trang chủ"]["engagementRate"], 537 / 540)

    def test_country_matrix_and_trend(self):
        m = self.ai.country_sources(D0, D0)["data"]["matrix"]
        self.assertEqual(m["Japan"], {"GEMINI": 66})
        self.assertEqual(m["Vietnam"]["CHATGPT"], 537)
        tr = self.ai.country_trend(D0, D0, top=1)["data"]
        self.assertEqual(tr["series"], ["Vietnam", "Quốc gia khác"])
        self.assertEqual(tr["points"][0]["Quốc gia khác"], 66)

    def test_organic_by_landing_is_separate_channel(self):
        org = self.ai.organic_by_landing(D0, D0)
        self.assertEqual(org["/blog-a/"], 50)
        self.assertEqual(org["/"], 6000)

    def test_landing_types_match_with_or_without_slash(self):
        from seo_report import urls as U
        from seo_report.web_ai import landing_types
        con = db.connect(":memory:")
        site = "https://matkinhtamduc.com/"
        db.replace_day(con, site, "2026-10-01", [{"keys": [site + "gong-kinh/"], "clicks": 1, "impressions": 1,
                                                  "ctr": 1, "position": 1}], [])
        U.assign_urls(con, site)
        self.assertEqual(landing_types(con, site, ["/gong-kinh", "(not set)", "/khong-co/"]),
                         {"/gong-kinh": "Danh mục", "(not set)": "Không xác định",
                          "/khong-co/": "Chưa có trong dữ liệu GSC"})

    def test_excel_export_sheets(self):
        import io
        from pathlib import Path
        from openpyxl import load_workbook
        from seo_report.ai_export import build_ai_excel
        from seo_report.config import Config
        f = Filters()
        state = {"start": D0, "end": D0, "source": None, "country": None, "landing": None,
                 "data": {"summary": self.ai.summary(D0, D0, f), "sources": self.ai.sources(D0, D0, f),
                          "trend": self.ai.trend(D0, D0, f),
                          "key_events": self.ai.key_events(D0, D0, ["booking_success"], f)},
                 "land": {"pages": self.ai.landing_pages(D0, D0, f), "organic": {"/": 6000}, "types": {"/": "Trang chủ"},
                          "segments": self.ai.segments(D0, D0, {"Trang chủ": ["/"]}, f)},
                 "geo": {"countries": self.ai.countries(D0, D0, f), "matrix": self.ai.country_sources(D0, D0, f)}}
        cfg = Config(site_url="https://x.vn/", db_path=Path("x"), client_secret=Path("x"), token_path=Path("x"),
                     ga4_property="properties/1")
        wb = load_workbook(io.BytesIO(build_ai_excel(state, cfg, "test")))
        self.assertEqual(wb.sheetnames, ["Thông tin", "Tổng quan", "Kênh", "Nguồn AI", "Nguồn chi tiết", "Key events",
                                         "Xu hướng", "Loại trang", "Trang đích", "Quốc gia", "Quốc gia × nguồn"])
        self.assertEqual(wb["Tổng quan"]["B2"].value, 609)
        self.assertEqual(wb["Tổng quan"]["B7"].number_format, "0.00%")  # session key event rate định dạng %


class PeriodAndCacheTest(unittest.TestCase):
    def test_previous_period_same_length(self):
        self.assertEqual(previous_period(date(2026, 10, 1), date(2026, 10, 28)),
                         (date(2026, 9, 3), date(2026, 9, 30)))
        self.assertEqual(previous_period(date(2026, 3, 1), date(2026, 3, 1)), (date(2026, 2, 28), date(2026, 2, 28)))

    def test_change_no_baseline(self):
        self.assertIsNone(change(5, 0))
        self.assertAlmostEqual(change(15, 10), 0.5)

    def test_iso_weeks_cross_year_and_partial(self):
        days = ["20251231", "20260101", "20260105"]
        pts = _bucket(days, {"AI": {d: 1 for d in days}}, "week", date(2025, 12, 31), date(2026, 1, 5))
        self.assertEqual([(p["week"], p["AI"], p["partial"]) for p in pts],
                         [("2025-12-29", 2, True), ("2026-01-05", 1, True)])

    def test_cache_hit_and_key_includes_filter(self):
        con = db.connect(":memory:")
        fake = FakeGA4(spec_dataset())
        client = GA4Client(fake, "properties/1", con, today=date(2026, 12, 1))
        a = AIAnalytics(client, today=date(2026, 12, 1))
        a.summary(D0, D0, compare=False)
        n = len(fake.calls)
        env = a.summary(D0, D0, compare=False)
        self.assertEqual(len(fake.calls), n)  # lần hai lấy từ cache
        self.assertTrue(env["sourceMeta"]["cached"])
        a.summary(D0, D0, Filters(country="Vietnam"), compare=False)
        self.assertGreater(len(fake.calls), n)  # đổi bộ lọc → truy vấn mới


if __name__ == "__main__":
    unittest.main()


class R4OpportunityTest(unittest.TestCase):
    SITE = "https://matkinhtamduc.com/"

    def setUp(self):
        from seo_report import urls as U
        self.con = db.connect(":memory:")
        s = self.SITE
        for day in ("2026-10-01", "2026-10-02"):
            db.replace_day(self.con, s, day, [
                {"keys": [s + "gong-kinh/"], "clicks": 10, "impressions": 3000, "ctr": 0, "position": 4},
                {"keys": [s + "gong-kinh/?utm_source=x"], "clicks": 0, "impressions": 1000, "ctr": 0, "position": 8},
                {"keys": [s + "cua-hang/"], "clicks": 1, "impressions": 20, "ctr": 0, "position": 12},
                {"keys": [s + "blog/bai-viet/"], "clicks": 5, "impressions": 900, "ctr": 0, "position": 6},
            ], [])
        U.assign_urls(self.con, s)
        self.ga4 = [
            {"landingPage": "/gong-kinh", "sessions": 30, "engagedSessions": 9, "keyEvents": 1,
             "sessionKeyEventRate": 0.03},
            {"landingPage": "/gong-kinh/", "sessions": 10, "engagedSessions": 1, "keyEvents": 0,
             "sessionKeyEventRate": 0.0},
            {"landingPage": "/cua-hang", "sessions": 8, "engagedSessions": 8, "keyEvents": 2,
             "sessionKeyEventRate": 0.25},
            {"landingPage": "/trang-moi", "sessions": 3, "engagedSessions": 3, "keyEvents": 0,
             "sessionKeyEventRate": 0.0},
            {"landingPage": "(not set)", "sessions": 49, "engagedSessions": 0, "keyEvents": 0,
             "sessionKeyEventRate": 0.0},
        ]

    def build(self):
        from seo_report import ai_opportunities as O
        return O.build(self.con, self.SITE, date(2026, 10, 1), date(2026, 10, 2), self.ga4, 0.80)

    def test_align_merges_paths_and_keeps_sources_apart(self):
        from seo_report import ai_opportunities as O
        rows = {r["path"]: r for r in self.build()["rows"]}
        g = rows["/gong-kinh/"]
        self.assertEqual(g["ga4_paths"], ["/gong-kinh", "/gong-kinh/"])
        self.assertEqual(g["ai_sessions"], 40)
        self.assertAlmostEqual(g["ai_engagement_rate"], 10 / 40)  # tính lại từ số đếm, không lấy trung bình rate
        self.assertIsNone(g["ai_session_key_event_rate"])  # không ghép được rate phiên có key event
        # GSC: ?utm gộp vào URL báo cáo; vị trí TB có trọng số impression
        self.assertEqual((g["gsc_clicks"], g["gsc_impressions"]), (20, 8000))
        self.assertAlmostEqual(g["gsc_position"], (4 * 6000 + 8 * 2000) / 8000)
        self.assertEqual(g["match"], O.MATCHED)
        self.assertEqual(rows["/trang-moi"]["match"], O.GA4_ONLY)
        self.assertEqual(rows["/trang-moi"]["type"], "Chưa có trong dữ liệu GSC")
        self.assertEqual(rows["/blog/bai-viet/"]["match"], O.GSC_ONLY)
        self.assertNotIn("(not set)", rows)

    def test_rules_fire_with_evidence(self):
        res = self.build()
        by = {(o["rule_id"], o["path"]) for o in res["opportunities"]}
        self.assertIn(("AI_REFERRAL_WITH_LOW_GSC_VISIBILITY", "/cua-hang/"), by)  # 8 phiên AI, 40 impression
        self.assertIn(("AI_ENGAGEMENT_WEAK", "/gong-kinh/"), by)  # 25% vs 80%, 40 phiên
        self.assertNotIn(("GSC_HIGH_IMPRESSIONS_LOW_AI_REFERRAL", "/gong-kinh/"), by)  # có 40 phiên AI
        from seo_report import ai_opportunities as O
        for o in res["opportunities"]:
            self.assertTrue(all(o.get(k) not in (None, "") for k in O.OPP_FIELDS), o)
        codes = {w["code"] for w in res["warnings"]}
        self.assertTrue({"TIMEZONE_MISMATCH", "URL_MATCH_AMBIGUOUS"} <= codes)

    def test_thresholds_saved_and_logged(self):
        from seo_report import ai_opportunities as O
        O.set_thresholds(self.con, self.SITE, {"high_gsc_min_impressions": 1000, "high_gsc_max_ai_sessions": 0},
                         "tester")
        self.assertEqual(O.get_thresholds(self.con)["high_gsc_min_impressions"], 1000)
        self.assertEqual(O.get_thresholds(self.con)["weak_min_sessions"], O.DEFAULTS["weak_min_sessions"])
        by = {(o["rule_id"], o["path"]) for o in self.build()["opportunities"]}
        self.assertIn(("GSC_HIGH_IMPRESSIONS_LOW_AI_REFERRAL", "/blog/bai-viet/"), by)  # 1.800 impression, 0 phiên AI
        self.assertEqual(self.con.execute("SELECT user FROM change_log WHERE action LIKE '%ngưỡng%'").fetchone()[0],
                         "tester")

    def test_no_engagement_rule_without_sample(self):
        from seo_report import ai_opportunities as O
        rows = O.align(self.con, self.SITE, [{"landingPage": "/cua-hang", "sessions": 5, "engagedSessions": 0,
                                              "keyEvents": 0, "sessionKeyEventRate": 0}], {})
        self.assertEqual([o for o in O.rules(rows, O.DEFAULTS, 0.8, "x") if o["rule_id"] == "AI_ENGAGEMENT_WEAK"], [])

    def test_excel_has_opportunity_sheets(self):
        import io
        from pathlib import Path
        from openpyxl import load_workbook
        from seo_report.ai_export import build_ai_excel
        from seo_report.config import Config
        cfg = Config(site_url=self.SITE, db_path=Path("x"), client_secret=Path("x"), token_path=Path("x"),
                     ga4_property="properties/1")
        state = {"start": D0, "end": D0, "source": None, "country": None, "landing": None, "data": None,
                 "opp": self.build()}
        wb = load_workbook(io.BytesIO(build_ai_excel(state, cfg, "test")))
        self.assertIn("Cơ hội - gợi ý", wb.sheetnames)
        self.assertEqual(wb["Cơ hội - URL"]["E1"].value, "GA4 AI · Phiên")
        self.assertEqual(wb["Cơ hội - URL"]["I1"].value, "GSC Web · Impression")

    def test_ga4_path_prefers_trailing_slash_url(self):
        """Site 308 /x → /x/: bản thiếu / trong GSC là URL cũ, phiên AI phải về bản có /."""
        from seo_report import urls as U
        db.replace_day(self.con, self.SITE, "2026-10-01", [
            {"keys": [self.SITE + "gong-kinh/"], "clicks": 10, "impressions": 3000, "ctr": 0, "position": 4},
            {"keys": [self.SITE + "gong-kinh"], "clicks": 0, "impressions": 2, "ctr": 0, "position": 9}], [])
        U.assign_urls(self.con, self.SITE)
        rows = {r["path"]: r for r in self.build()["rows"]}
        self.assertEqual(rows["/gong-kinh/"]["ai_sessions"], 40)
        self.assertEqual(rows["/gong-kinh"]["ai_sessions"], 0)  # vẫn hiện số GSC của URL cũ, không nhận phiên AI


class R5QualityTest(unittest.TestCase):
    def setUp(self):
        data = spec_dataset()
        data += sess("chat.mistral.ai", "referral", "Referral", 4)  # có dáng AI nhưng chưa có trong danh sách
        data[600]["landingPage"] = "(not set)"
        self.ai = make(data)

    def envs(self):
        f = Filters()
        return (self.ai.summary(D0, D0, f, compare=False), self.ai.sources(D0, D0, f, compare=False),
                self.ai.landing_pages(D0, D0, f, compare=False, limit=1000),
                self.ai.countries(D0, D0, f, compare=False))

    def test_totals_match_and_rates_in_range(self):
        from seo_report import ai_quality as Q
        checks = {c["code"]: c for c in Q.ga4_checks(*self.envs())}
        self.assertEqual(checks["GA4_TOTAL_DISCREPANCY"]["status"], Q.OK)
        self.assertEqual(checks["RATE_RANGE"]["status"], Q.OK)
        self.assertIn("1 phiên AI không xác định", checks["LANDING_NOT_SET"]["title"])

    def test_discrepancy_and_bad_rate_flagged(self):
        from seo_report import ai_quality as Q
        summ, src, land, geo = self.envs()
        land["data"]["rows"] = land["data"]["rows"][1:]          # mất dòng lớn nhất (giống bị cắt dòng)
        src["data"]["sources"][0]["engagementRate"] = 73.08      # lỗi kiểu nhân 100 hai lần
        checks = {c["code"]: c for c in Q.ga4_checks(summ, src, land, geo)}
        self.assertEqual(checks["GA4_TOTAL_DISCREPANCY"]["status"], Q.WARN)
        self.assertEqual(checks["RATE_RANGE"]["status"], Q.ERROR)
        self.assertEqual(Q.summarize(list(checks.values())), Q.ERROR)

    def test_api_error_is_error_not_zero(self):
        from seo_report import ai_quality as Q
        summ, src, land, geo = self.envs()
        src = {**src, "status": "ERROR", "warnings": [{"code": "API_PERMISSION_DENIED", "message": "403"}]}
        checks = Q.ga4_checks(summ, src, land, geo)
        self.assertEqual([c["status"] for c in checks], [Q.ERROR])

    def test_inventory_and_candidates(self):
        from seo_report import ai_quality as Q
        inv = Q.source_inventory(self.envs()[1])
        self.assertEqual(inv[C.NATIVE], 609)
        cands = {c["source"]: c for c in Q.candidate_sources(self.ai, D0, D0)}
        self.assertEqual(cands["chatgpt"]["reason"][:14], "Medium trả phí")   # quảng cáo, theo quyết định R0
        self.assertEqual(cands["chat.mistral.ai"]["sessions"], 4)
        self.assertNotIn("chatgpt.com", cands)                              # đã tính là AI thì không liệt kê

    def test_gsc_status_gaps_and_generative_ai_unavailable(self):
        from seo_report import ai_quality as Q
        con = db.connect(":memory:")
        site = "https://x.vn/"
        for d in ("2026-10-25", "2026-10-27"):
            db.replace_day(con, site, d, [], [])
        checks = {c["code"]: c for c in Q.gsc_status(con, site, date(2026, 10, 30))}
        self.assertEqual(checks["GSC_FRESHNESS"]["status"], Q.OK)            # trễ 3 ngày là bình thường
        self.assertIn("2026-10-26", checks["GSC_GAPS"]["detail"])
        self.assertIn("Not available", checks["NO_GENERATIVE_AI_SOURCE"]["title"])
        self.assertEqual(Q.gsc_status(db.connect(":memory:"), site, date(2026, 10, 30))[0]["status"], Q.ERROR)

    def test_cache_isolated_per_property(self):
        con = db.connect(":memory:")
        a = GA4Client(FakeGA4(sess("chatgpt.com", "x", "AI Assistant", 5)), "properties/1", con)
        b = GA4Client(FakeGA4(sess("chatgpt.com", "x", "AI Assistant", 2)), "properties/2", con)
        self.assertEqual(a.run([], ["sessions"], D0, D0).rows[0]["sessions"], 5)
        self.assertEqual(b.run([], ["sessions"], D0, D0).rows[0]["sessions"], 2)  # không lấy cache của property 1

    def test_small_base_hides_percent(self):
        from seo_report import ai_format as F
        self.assertEqual(F.delta_pct(16.0, 1), "kỳ trước 1")
        self.assertEqual(F.delta_pct(0.25, 40), "+25,0%")
        self.assertEqual(F.delta_pct(None), "")

    def test_url_params_roundtrip(self):
        from seo_report.web_ai import params_from_state, state_from_params
        today = date(2026, 10, 30)
        st = state_from_params({"preset": "custom", "from": "2026-09-01", "to": "2026-09-30", "compare": "0",
                                "source": "GEMINI", "tab": "co-hoi"}, today)
        self.assertEqual((st["start"], st["end"], st["compare"], st["source"], st["tab"]),
                         (date(2026, 9, 1), date(2026, 9, 30), False, "GEMINI", "Cơ hội SEO/GEO"))
        self.assertEqual(params_from_state(st, st["tab"]), {"preset": "custom", "from": "2026-09-01",
                                                            "to": "2026-09-30", "compare": "0", "source": "GEMINI",
                                                            "tab": "co-hoi"})
        bad = state_from_params({"preset": "custom", "from": "2026-09-30", "to": "x", "tab": "?"}, today)
        self.assertEqual((bad["preset"], bad["tab"], bad["compare"]), ("28", "Tổng quan", True))


def gai_zip(days: dict, pages=None, filters=None, countries=None, extra=None) -> bytes:
    """File export giống GSC thật (zip 5 CSV, không có dòng trống cuối)."""
    import io as _io
    import zipfile
    files = {
        "Chart.csv": "Date,Impressions\n" + "\n".join(f"{d},{v}" for d, v in days.items()),
        "Pages.csv": "Top pages,Impressions\n" + "\n".join(f"{p},{v}" for p, v in (pages or {}).items()),
        "Countries.csv": "Country,Impressions\n" + "\n".join(f"{k},{v}" for k, v in
                                                             (countries or {"Vietnam": 10}).items()),
        "Devices.csv": "Device,Impressions\nMobile,7\nDesktop,3",
        "Filters.csv": "Filter,Value\n" + "\n".join(f"{k},{v}" for k, v in
                                                     (filters or {"Search type": "Web (text-based)",
                                                                  "Date": "Last 3 months"}).items()),
        **(extra or {}),
    }
    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, text in files.items():
            z.writestr(name, text)
    return buf.getvalue()


class R6GenerativeAIImportTest(unittest.TestCase):
    SITE = "https://matkinhtamduc.com/"

    def setUp(self):
        from seo_report import urls as U
        self.con = db.connect(":memory:")
        db.replace_day(self.con, self.SITE, "2026-10-01", [
            {"keys": [self.SITE + "gong-kinh/"], "clicks": 1, "impressions": 100, "ctr": 0, "position": 3}], [])
        U.assign_urls(self.con, self.SITE)
        for d in range(1, 15):
            self.con.execute("INSERT OR REPLACE INTO gsc_site_daily VALUES (?, ?, 0, 1000, 5)",
                             (self.SITE, f"2026-10-{d:02d}"))

    def days(self, a, b, v=100):
        return {f"2026-10-{d:02d}": v for d in range(a, b + 1)}

    def imp(self, content, name="x.zip"):
        from seo_report import ai_gsc_import as G
        return G.save(self.con, self.SITE, G.parse(name, content, self.SITE), name, "tester")

    def test_parse_real_shape_zero_is_unknown(self):
        from seo_report import ai_gsc_import as G
        days = self.days(1, 7)
        days["2026-10-03"] = 0
        p = G.parse("x.zip", gai_zip(days, {self.SITE + "gong-kinh/?page=2": 5, self.SITE: 0}), self.SITE)
        self.assertEqual((p["period_start"], p["period_end"], p["search_type"]), ("2026-10-01", "2026-10-07", "text"))
        self.assertIsNone(p["daily"]["2026-10-03"])   # GSC hiện "~": không phải 0
        self.assertIsNone(p["pages"][self.SITE])
        self.assertEqual(p["stats"]["unknown_zero"]["chart"], 1)
        self.assertEqual(p["stats"]["total_chart"], 600)

    def test_vietnamese_headers_and_xlsx(self):
        import io as _io
        from openpyxl import Workbook
        from seo_report import ai_gsc_import as G
        wb = Workbook()
        ws = wb.active
        ws.append(["Ngày", "Số lượt hiển thị"])
        ws.append(["2026-10-01", 50])
        ws.append(["2026-10-02", 4630.0])  # Excel hay lưu số dạng float: không được đọc thành 46300
        ws2 = wb.create_sheet()
        ws2.append(["Trang hàng đầu", "Số lượt hiển thị"])
        ws2.append([self.SITE + "gong-kinh/", 20])
        ws3 = wb.create_sheet()
        ws3.append(["Bộ lọc", "Giá trị"])
        ws3.append(["Loại tìm kiếm", "Web (đa phương thức)"])
        buf = _io.BytesIO()
        wb.save(buf)
        p = G.parse("x.xlsx", buf.getvalue(), self.SITE)
        self.assertEqual((p["daily"], p["pages"], p["search_type"]),
                         ({"2026-10-01": 50, "2026-10-02": 4630}, {self.SITE + "gong-kinh/": 20}, "multimodal"))

    def test_rejects_bad_files(self):
        from seo_report import ai_gsc_import as G
        bad = [gai_zip(self.days(1, 2), filters={"Search type": "Web (text-based)", "Page": "+gong-kinh"}),
               gai_zip(self.days(1, 2), {"https://other.vn/a/": 3}),
               gai_zip({}, extra={"Chart.csv": "Date,Impressions"})]  # không có ngày nào
        for content in bad:
            with self.assertRaises(G.ImportError_):
                G.parse("x.zip", content, self.SITE)
        with self.assertRaises(G.ImportError_):
            G.parse("x.csv", b"Date,Impressions", self.SITE)

    def test_overlap_latest_wins_and_delete_restores(self):
        from seo_report import ai_gsc_import as G
        first = self.imp(gai_zip(self.days(1, 14, 100)))
        second = self.imp(gai_zip(self.days(8, 14, 150)))
        d = G.daily(self.con, self.SITE, date(2026, 10, 1), date(2026, 10, 14))
        self.assertEqual((d["2026-10-07"], d["2026-10-08"]), (100, 150))  # không cộng dồn hai lần nhập
        tot = G.period_total(self.con, self.SITE, date(2026, 10, 1), date(2026, 10, 14))
        self.assertEqual((tot["status"], tot["impressions"]), ("MANUAL_IMPORT", 7 * 100 + 7 * 150))
        G.delete(self.con, self.SITE, second, "tester")
        self.assertEqual(G.daily(self.con, self.SITE, date(2026, 10, 8), date(2026, 10, 8)), {"2026-10-08": 100})
        self.assertEqual([i["id"] for i in G.imports(self.con, self.SITE)], [first])
        self.assertEqual(len(G.imports(self.con, self.SITE, include_deleted=True)), 2)

    def test_status_not_connected_and_partial(self):
        from seo_report import ai_gsc_import as G
        self.assertEqual(G.period_total(self.con, self.SITE, date(2026, 10, 1), date(2026, 10, 7))["status"],
                         "NOT_CONNECTED")
        self.imp(gai_zip(self.days(1, 5)))
        t = G.period_total(self.con, self.SITE, date(2026, 10, 1), date(2026, 10, 7))
        self.assertEqual((t["status"], t["impressions"], t["covered"]), ("PARTIAL", None, 5))  # không hiện tổng thiếu

    def test_page_matched_to_report_url_raw_kept(self):
        from seo_report import ai_gsc_import as G
        iid = self.imp(gai_zip(self.days(1, 2), {self.SITE + "gong-kinh/?utm_source=x": 9}))
        row = G.breakdown(self.con, iid, "page")[0]
        self.assertEqual(row["key"], self.SITE + "gong-kinh/?utm_source=x")
        self.assertIsNotNone(row["url_id"])

    def test_weekly_block_and_quality_status(self):
        from seo_report import ai_quality as Q
        from seo_report.report import aio_block
        self.assertIsNone(aio_block(self.con, self.SITE, (date(2026, 10, 5), date(2026, 10, 11)),
                                    (date(2026, 9, 28), date(2026, 10, 4))))
        self.assertEqual(Q.generative_ai_status(self.con, self.SITE, date(2026, 10, 30))["code"],
                         "NO_GENERATIVE_AI_SOURCE")
        self.imp(gai_zip(self.days(1, 14, 50)))
        b = aio_block(self.con, self.SITE, (date(2026, 10, 5), date(2026, 10, 11)),
                      (date(2026, 9, 28), date(2026, 10, 4)))
        self.assertEqual((b["impressions"], b["prev"], b["share"]), (350, None, 350 / 7000))
        st = Q.generative_ai_status(self.con, self.SITE, date(2026, 10, 30))
        self.assertEqual(st["status"], Q.WARN)  # file tới 14/10, đã 16 ngày
