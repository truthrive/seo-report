"""Volume keyword SEO (DataForSEO), lưu Google Drive (drive.file), bộ slide báo cáo."""
import io
import unittest
from datetime import date
from pathlib import Path

from seo_report import db, drive as D, rank as R, report, volume as V
from seo_report import urls as U
from seo_report.config import Config

SITE = "https://matkinhtamduc.com/"
A = SITE + "gong-kinh/"
TODAY = date(2026, 10, 30)


def cfg_with_secret(tmp: Path) -> Config:
    secret = tmp / "dfs.toml"
    secret.write_text('login = "x"\npassword = "y"\n', encoding="utf-8")
    return Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"),
                  dataforseo_secret=secret)


class VolumeTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = Path(tempfile.mkdtemp())
        self.cfg = cfg_with_secret(self.tmp)
        self.con = db.connect(":memory:")
        R.add_keywords(self.con, SITE, self.cfg.rank, TODAY, [("mắt kính", None), ("gọng kính", A)], "t")
        self.calls = []

    def fetcher(self, login, pw, keywords, loc, lang):
        self.calls.append(keywords)
        return {"cost": 0.075, "result": [{"keyword": "mắt kính", "search_volume": 33100, "competition": "MEDIUM"}]}

    def test_refresh_once_per_month_and_unknown_kept(self):
        n = V.refresh(self.con, SITE, self.cfg, TODAY, TODAY, fetcher=self.fetcher, log=lambda m: None)
        self.assertEqual(n, 2)
        vols = {k["keyword"]: V.latest(self.con).get(k["id"]) for k in R.keywords(self.con, SITE, TODAY)}
        self.assertEqual(vols["mắt kính"]["volume"], 33100)
        self.assertIsNone(vols["gọng kính"]["volume"])  # Google Ads không trả: lưu None, không hỏi lại trong tháng
        V.refresh(self.con, SITE, self.cfg, TODAY, TODAY, fetcher=self.fetcher, log=lambda m: None)
        self.assertEqual(len(self.calls), 1)
        V.refresh(self.con, SITE, self.cfg, TODAY, date(2026, 11, 2), fetcher=self.fetcher, log=lambda m: None)
        self.assertEqual(len(self.calls), 2)  # tháng mới thì lấy lại

    def test_cost_counts_in_rank_budget(self):
        V.refresh(self.con, SITE, self.cfg, TODAY, TODAY, fetcher=self.fetcher, log=lambda m: None)
        # Chi phí volume ghi theo giờ thật (run_at), nên so với tháng hiện tại chứ không phải TODAY cố định
        self.assertAlmostEqual(R.spent_this_month(self.con, date.today()) - R.spent_this_month(self.con, date(2000, 1, 1)), 0.075)

    def test_budget_blocks_fetch(self):
        self.cfg.rank.monthly_budget_usd = 0.01
        logs = []
        self.assertEqual(V.refresh(self.con, SITE, self.cfg, TODAY, TODAY, fetcher=self.fetcher, log=logs.append), 0)
        self.assertEqual(self.calls, [])
        self.assertIn("ngân sách", logs[0])


class FakeReq:
    def __init__(self, fn):
        self.fn = fn

    def execute(self):
        return self.fn()


class FakeDrive:
    """Drive giả: chỉ nhớ file do 'tool' tạo (giống quyền drive.file)."""

    def __init__(self):
        self.files_, self.n = {}, 0

    def files(self):
        return self

    def create(self, body, fields=None, media_body=None):
        def go():
            self.n += 1
            fid = f"id{self.n}"
            self.files_[fid] = {**body, "id": fid, "trashed": False, "media": media_body}
            return {"id": fid, "name": body["name"], "webViewLink": f"https://drive/{fid}"}
        return FakeReq(go)

    def get(self, fileId, fields=None):
        def go():
            if fileId not in self.files_:
                from googleapiclient.errors import HttpError
                raise HttpError(type("R", (), {"status": 404, "reason": "nf"})(), b"{}")
            return self.files_[fileId]
        return FakeReq(go)

    def list(self, q, fields=None, orderBy=None, pageSize=None):
        parent = q.split("'")[1]
        return FakeReq(lambda: {"files": [{"id": f["id"], "name": f["name"]} for f in self.files_.values()
                                          if parent in f.get("parents", []) and f.get("mimeType") == D.FOLDER]})

    def export(self, fileId, mimeType):
        return FakeReq(lambda: b"%PDF-fake")


class DriveTest(unittest.TestCase):
    def setUp(self):
        self.con = db.connect(":memory:")
        self.svc = FakeDrive()

    def test_root_folder_created_once_and_recreated_if_deleted(self):
        root = D.root_folder(self.con, self.svc, SITE)
        self.assertEqual(self.svc.files_[root]["name"], "SEO Report – matkinhtamduc.com")
        self.assertEqual(D.root_folder(self.con, self.svc, SITE), root)
        del self.svc.files_[root]
        self.assertNotEqual(D.root_folder(self.con, self.svc, SITE), root)

    def test_target_folder_and_convert_upload(self):
        root = D.root_folder(self.con, self.svc, SITE)
        sub = D.create_folder(self.svc, "Báo cáo tháng", root)
        self.assertEqual([f["name"] for f in D.list_folders(self.svc, root)], ["Báo cáo tháng"])
        self.assertEqual(D.target_folder(self.con, self.svc, SITE)[0], root)  # mặc định: folder gốc
        D.set_target_folder(self.con, sub)
        self.assertEqual(D.target_folder(self.con, self.svc, SITE), (sub, "Báo cáo tháng"))
        f = D.upload(self.svc, b"xlsx", "bao-cao", D.XLSX, sub, D.SHEET)
        self.assertEqual(self.svc.files_[f["id"]]["mimeType"], D.SHEET)  # Drive tự chuyển sang Google Sheets
        self.assertEqual(self.svc.files_[f["id"]]["parents"], [sub])
        self.assertEqual(D.export_pdf(self.svc, f["id"]), b"%PDF-fake")

    def test_explain_missing_scope(self):
        self.assertIn("login", D.explain(SystemExit("Tài khoản chưa cấp quyền Google Drive")))

    def test_login_requests_drive_scope(self):
        from seo_report import auth
        self.assertIn(auth.DRIVE_SCOPE, auth.SCOPES)
        self.assertNotIn(auth.DRIVE_SCOPE, auth.ALL_SCOPES)  # chức năng GA4 không đòi quyền Drive


class SlidesTest(unittest.TestCase):
    def test_week_and_month_decks(self):
        from pptx import Presentation

        from seo_report.report_slides import build_report_pptx, slides_filename
        from tests.test_report import MonthlyReportTest
        t = MonthlyReportTest()
        t.setUp()
        report.save_note(t.con, SITE, report.note_key(t.rep), "Tháng tốt.", "t")
        rep = report.build_month(t.con, t.cfg, t.gid, 0, with_ai=False)
        prs = Presentation(io.BytesIO(build_report_pptx(rep, "t")))
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
        self.assertIn("Báo cáo SEO & GEO tháng 10/2026", texts)
        self.assertIn("Tháng tốt.", texts)
        self.assertNotIn("Chưa có kế hoạch.", texts)  # không để chữ "chưa có" trong bản gửi đi
        self.assertEqual(len(prs.slides), 8)  # bìa, kết luận, SEO, GEO, thứ hạng, 3 phụ lục
        for title in ("Kết luận", "Kết quả SEO", "Kết quả GEO", "Thứ hạng từ khóa theo dõi"):
            self.assertIn(title, texts)
        self.assertFalse(any("CTR" in t for t in texts))
        self.assertEqual(slides_filename(rep), "bao-cao-seo-thang_matkinhtamduc_2026-10")
        week = report.build(t.con, t.cfg, t.gid, 0, with_ai=False)
        self.assertEqual(len(Presentation(io.BytesIO(build_report_pptx(week, "t"))).slides), 4)  # không nhận xét, không điểm chú ý: bỏ slide kết luận
