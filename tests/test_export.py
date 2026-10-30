import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from seo_report import db, export
from seo_report import urls as U
from seo_report.config import Config

SITE = "https://matkinhtamduc.com/"
A = "https://matkinhtamduc.com/gong-kinh/"
B = "https://matkinhtamduc.com/"


def page(url, c, i):
    return {"keys": [url], "clicks": c, "impressions": i, "ctr": 0, "position": 4.0}


def q(url, query, c, i):
    return {"keys": [url, query], "clicks": c, "impressions": i, "ctr": 0, "position": 4.0}


class ExportTest(unittest.TestCase):
    def setUp(self):
        self.con = db.connect(":memory:")
        # Tuần 19–25/10 và tuần trước 12–18/10; dữ liệu đến 28/10.
        db.replace_day(self.con, SITE, "2026-10-20", [page(A, 30, 300), page(B, 10, 200)],
                       [q(A, "gọng kính", 20, 200), q(A, "mắt kính", 2, 100), q(B, "mắt kính", 8, 150)])
        db.replace_day(self.con, SITE, "2026-10-13", [page(A, 20, 250)], [q(A, "gọng kính", 15, 180)])
        db.replace_day(self.con, SITE, "2026-10-28", [], [])
        U.assign_urls(self.con, SITE)
        self.gid = U.ensure_group(self.con, SITE)
        U.add_to_group(self.con, SITE, self.gid, A, "test")
        self.cfg = Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"))
        data = export.build_data_export(self.con, self.cfg, self.gid, 0, "test")
        self.wb = load_workbook(io.BytesIO(data))

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames,
                         ["Thông tin", "Tổng quan", "URL", "URL_Keyword", "Cannibalization", "Rank_History",
                          "Rank_History_Long", "URL_Daily"])

    def test_url_sheet_numbers(self):
        ws = self.wb["URL"]
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(row["URL"], A)
        self.assertEqual((row["Click"], row["Click kỳ trước"], row["Chênh click"]), (30, 20, 10))
        self.assertAlmostEqual(row["% click"], 0.5)
        self.assertEqual(row["Head KW (28 ngày)"], "gọng kính")
        self.assertEqual(row["Click nhìn thấy theo keyword"], 22)

    def test_keyword_roles_and_cannibal(self):
        ws = self.wb["URL_Keyword"]
        roles = {r[1]: (r[2], r[3]) for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(roles["gọng kính"], ("Head", None))  # ô trống
        self.assertEqual(roles["mắt kính"], ("Bị loại", "URL khác mạnh hơn"))
        cannibal = [r[:4] for r in self.wb["Cannibalization"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(cannibal, [("mắt kính", B, "Không", "Có"), ("mắt kính", A, "Có", None)])

    def test_total_row_and_filters(self):
        ws = self.wb["Tổng quan"]
        self.assertEqual(ws.cell(ws.max_row, 1).value, "Tất cả")
        self.assertEqual(ws.cell(ws.max_row, 3).value, 30)
        self.assertEqual(self.wb["URL_Daily"].freeze_panes, "A2")
        self.assertTrue(self.wb["URL"].auto_filter.ref)


if __name__ == "__main__":
    unittest.main()
