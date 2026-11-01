import sqlite3
import unittest
from datetime import date, timedelta

from seo_report import kpi as KP
from seo_report import rank as R

SITE = "https://x.vn/"
TODAY = date(2027, 1, 20)
ROWS = [
    ["KPI cam kết SEO"],
    [],
    ["Tên dự án", "SEO x.vn"],
    ["Ngày bắt đầu", "01/11/2026"],
    [],
    ["Tháng thứ dự án", 1, 2, 3],
    ["Giai đoạn", "Review tối ưu nền tảng", "", "Thúc đẩy"],
    ["Số lượng từ khóa vào top 3 (cam kết)", "-%", "-", "50%"],
    ["Số lượng từ khóa vào top 5 (cam kết)", "-", "", 0.5],
    ["Số lượng từ khóa vào top 10 (cam kết)", "", "25%", 75],
    ["Traffic Session Organic Search (GA4)", "39,000", 39000, "59.348"],
]


class ParseTest(unittest.TestCase):
    def test_parse_kpi(self):
        d = KP.parse_kpi(ROWS, TODAY)
        self.assertEqual((d["project"], d["start"]), ("SEO x.vn", "2026-11-01"))
        m1, m2, m3 = d["months"]
        self.assertEqual((m1["top3"], m1["top10"], m1["sessions"]), (None, None, 39000))
        self.assertEqual(m2["phase"], "Review tối ưu nền tảng")  # ô gộp: lấy giai đoạn bên trái
        self.assertEqual((m2["top10"], m3["top3"], m3["top5"], m3["top10"]), (0.25, 0.5, 0.5, 0.75))
        self.assertEqual(m3["sessions"], 59348)

    def test_more_months_added_later(self):
        rows = [r[:] for r in ROWS]
        rows[5] = ["Tháng thứ dự án"] + list(range(1, 10))           # dự án kéo dài thành 9 tháng
        rows[9] = ["Số lượng từ khóa vào top 10 (cam kết)"] + ["-"] * 8 + ["70%"]
        d = KP.parse_kpi(rows, TODAY)
        self.assertEqual([m["n"] for m in d["months"]], list(range(1, 10)))
        self.assertEqual(d["months"][8]["top10"], 0.7)

    def test_keywords_dedup_and_header(self):
        kws = KP.parse_keywords([["Từ khóa", "URL mục tiêu"], ["kính A", "https://x.vn/a/"], ["Kính A", ""],
                                 ["kính B"], ["", "x"]])
        self.assertEqual(kws, [{"keyword": "kính A", "url": "https://x.vn/a/", "target": None},
                               {"keyword": "kính B", "url": "", "target": None}])

    def test_keywords_by_header_with_top_target(self):
        kws = KP.parse_keywords([["Ghi chú", "Từ khóa", "Top target", "URL mục tiêu"],
                                 ["", "kính A", "Top 3", "https://x.vn/a/"], ["x", "kính B", 10, ""],
                                 ["", "kính C", "", ""]])
        self.assertEqual([(k["keyword"], k["target"], k["url"]) for k in kws],
                         [("kính A", 3, "https://x.vn/a/"), ("kính B", 10, ""), ("kính C", None, "")])

    def test_month_window(self):
        self.assertEqual(KP.month_window(date(2026, 11, 1), 1), (date(2026, 11, 1), date(2026, 11, 30)))
        self.assertEqual(KP.month_window(date(2026, 11, 1), 3), (date(2027, 1, 1), date(2027, 1, 31)))
        self.assertEqual(KP.month_window(date(2026, 1, 31), 2), (date(2026, 2, 28), date(2026, 3, 30)))


class ProgressTest(unittest.TestCase):
    def setUp(self):
        self.con = sqlite3.connect(":memory:")
        R.ensure_tables(self.con)
        self.ids = {}
        for kw in ("a", "b", "c", "d"):
            cur = self.con.execute("INSERT INTO rank_keywords (site, keyword, url_id, source, added_by, added_at) "
                                   "VALUES (?, ?, NULL, 'kpi', 't', '2026-11-01')", (SITE, kw))
            self.ids[kw] = cur.lastrowid
        data = KP.parse_kpi(ROWS, TODAY)
        data["keywords"] = [{"keyword": k, "url": "", "target": t} for k, t in
                            (("a", 3), ("b", 3), ("c", 10), ("d", None))]
        KP.save(self.con, data, date(2026, 11, 1))

    def tearDown(self):
        self.con.close()

    def check(self, kw, day, rank):
        self.con.execute("INSERT INTO rank_checks VALUES (?, ?, 30, ?, NULL, NULL, 0, ?)",
                         (self.ids[kw], day, rank, day))

    def test_last_check_in_month_and_sessions(self):
        self.check("a", "2027-01-05", 15)
        self.check("a", "2027-01-15", 2)   # lần cuối trong tháng 3 → top 3
        self.check("b", "2027-01-15", 4)
        self.check("c", "2026-12-20", 1)   # check ở tháng 2, tháng 3 chưa check → chưa đạt
        sessions = lambda s, e: {(s + timedelta(days=i)).isoformat(): 1000 for i in range((e - s).days + 1)}
        p = KP.progress(self.con, SITE, TODAY, sessions)
        self.assertEqual(p["current"], 3)
        m3 = p["months"][2]
        self.assertEqual((m3["rank"]["top3"], m3["rank"]["top5"], m3["rank"]["top10"], m3["rank"]["checked"]),
                         (1, 2, 2, 2))
        ch = m3["checks"]
        self.assertEqual((ch["top3"]["need"], ch["top3"]["gap"], ch["top3"]["ok"]), (2, 1, False))
        self.assertEqual((ch["top5"]["ok"], ch["top10"]["gap"]), (True, 1))
        self.assertEqual(m3["actual_sessions"], 19000)    # 01/01 → 19/01 (hôm qua)
        self.assertEqual(m3["forecast"], 31000)           # 1000/ngày × 31 ngày
        self.assertFalse(ch["sessions"]["ok"])            # dự báo 31.000 < cam kết 59.348
        self.assertEqual((m3["rank"]["target_set"], m3["rank"]["target_ok"]), (3, 1))  # a đạt top 3, b hạng 4
        items = {x["keyword"]: x["ok"] for x in m3["rank"]["items"]}
        self.assertEqual(items, {"a": True, "b": False, "c": False, "d": None})
        m1 = p["months"][0]
        self.assertEqual(m1["status"], "closed")
        self.assertIsNone(m1["checks"]["top10"]["ok"])    # tháng không cam kết: không chấm

    def test_closed_month_uses_version_at_month_end(self):
        # Điều chỉnh KPI giữa tháng 3: tháng 2 (đã chốt) vẫn chấm theo bản cũ
        data = KP.parse_kpi(ROWS, TODAY)
        data["months"][1]["top10"] = 0.9
        data["keywords"] = [{"keyword": k, "url": ""} for k in ("a", "b", "c", "d")]
        self.assertTrue(KP.save(self.con, data, date(2027, 1, 10)))
        self.assertFalse(KP.save(self.con, data, date(2027, 1, 11)))  # không đổi thì không lưu thêm
        p = KP.progress(self.con, SITE, TODAY, None)
        self.assertEqual(p["months"][1]["top10"], 0.25)
        self.assertIsNone(p["months"][2]["actual_sessions"])          # chưa kết nối GA4: không hiện 0

    def test_template_roundtrip(self):
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(KP.template_xlsx("SEO x.vn", [("kính A", "https://x.vn/a/")])))
        rows = [list(r) for r in wb[KP.KPI_TAB].iter_rows(values_only=True)]
        d = KP.parse_kpi(rows, TODAY)
        self.assertEqual([m["top10"] for m in d["months"]], [None, None, 0.34, 0.40, 0.49, 0.63])
        self.assertEqual(d["months"][5]["sessions"], 71228)
        self.assertIsNone(d["start"])
        kws = KP.parse_keywords([list(r) for r in wb[KP.KW_TAB].iter_rows(values_only=True)])
        self.assertEqual(kws, [{"keyword": "kính A", "url": "https://x.vn/a/", "target": None}])


if __name__ == "__main__":
    unittest.main()
