import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from seo_report import db, report
from seo_report import rank as R
from seo_report import urls as U
from seo_report.config import Config
from seo_report.report_excel import build_report_excel

SITE = "https://matkinhtamduc.com/"
A = "https://matkinhtamduc.com/gong-kinh/"
B = "https://matkinhtamduc.com/blog-a/"
C = "https://matkinhtamduc.com/blog-nho/"
OTHER = "https://matkinhtamduc.com/ngoai-nhom/"


def page(url, c):
    return {"keys": [url], "clicks": c, "impressions": c * 10, "ctr": 0, "position": 5.0}


class ReportTest(unittest.TestCase):
    def setUp(self):
        self.con = db.connect(":memory:")
        # Tuần này 19–25/10, tuần trước 12–18/10; dữ liệu đến 28/10.
        db.replace_day(self.con, SITE, "2026-10-20", [page(A, 50), page(B, 5), page(C, 12), page(OTHER, 100)], [])
        db.replace_day(self.con, SITE, "2026-10-13", [page(A, 30), page(B, 25), page(C, 4), page(OTHER, 80)], [])
        db.replace_day(self.con, SITE, "2026-10-28", [], [])
        # Tổng theo property khác tổng cộng từ page (167/139): báo cáo phải dùng số này.
        db.save_site_totals(self.con, SITE, [
            {"keys": [d], "clicks": c, "impressions": c * 8, "ctr": 0, "position": 5.0}
            for d, c in [("2026-10-20", 160), ("2026-10-13", 130)]])
        U.assign_urls(self.con, SITE)
        self.gid = U.ensure_group(self.con, SITE)
        U.add_to_group(self.con, SITE, self.gid, f"{A}\n{B}\n{C}", "test")
        self.cfg = Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"))
        R.add_keywords(self.con, SITE, self.cfg.rank, report.date(2026, 10, 28), [("gọng kính", A)], "test")
        kid = R.keywords(self.con, SITE, report.date(2026, 10, 28))[0]["id"]
        for d, rank, url in [("2026-10-14", 5, A), ("2026-10-21", 2, B)]:
            self.con.execute("INSERT INTO rank_checks VALUES (?, ?, 30, ?, ?, NULL, 0.0018, 'x')", (kid, d, rank, url))
        self.rep = report.build(self.con, self.cfg, self.gid, 0)

    def test_kpis_site_vs_group(self):
        self.assertEqual((self.rep.site_kpi.clicks, self.rep.site_prev.clicks), (160, 130))
        self.assertEqual((self.rep.group_kpi.clicks, self.rep.group_prev.clicks), (67, 59))
        self.assertEqual(self.rep.week, (report.date(2026, 10, 19), report.date(2026, 10, 25)))

    def test_movers_skip_small_urls(self):
        self.assertEqual([(x["url"], x["diff"]) for x in self.rep.winners], [(A, 20)])
        self.assertEqual([(x["url"], x["diff"]) for x in self.rep.losers], [(B, -20)])  # C chỉ 12/4 click: bỏ qua

    def test_rank_week_and_wrong_url(self):
        r = self.rep.ranks[0]
        self.assertEqual((r["keyword"], r["rank"], r["prev"], r["ranking_url"]), ("gọng kính", 2, 5, B))
        self.assertEqual(self.rep.rank_buckets["Top 3"], [1, 0])
        self.assertEqual(self.rep.rank_buckets["Top 4–10"], [0, 1])

    def test_note_and_excel(self):
        report.save_note(self.con, SITE, report.date(2026, 10, 19), "Tuần ổn định.", "test")
        rep = report.build(self.con, self.cfg, self.gid, 0)
        ws = load_workbook(io.BytesIO(build_report_excel(rep, "test")))["Báo cáo"]
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertIn("Báo cáo SEO & GEO tuần 19/10/2026 – 25/10/2026", texts)
        self.assertIn("Tuần ổn định.", texts)
        self.assertIn("Sai URL mục tiêu", texts)
        self.assertEqual(len(ws._charts), 1)


class ReportAiBlockTest(unittest.TestCase):
    def test_no_ga4_means_no_block(self):
        cfg = Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"))
        self.assertIsNone(report.ai_block(cfg, None, (report.date(2026, 10, 19), report.date(2026, 10, 25))))

    def test_excel_ai_section_and_error(self):
        con = db.connect(":memory:")
        db.replace_day(con, SITE, "2026-10-28", [], [])
        gid = U.ensure_group(con, SITE)
        cfg = Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"))
        rep = report.build(con, cfg, gid, 0)
        rep.ai = {"kpi": {"sessions": 65, "engagementRate": 0.6, "keyEvents": 4, "sessionKeyEventRate": 0.0154},
                  "deltas": {"sessions": 0.91, "engagementRate": -0.02}, "share": {"ofAllSessions": 0.004},
                  "paid": {"sessions": 33}, "sources": [{"source": "CHATGPT", "sessions": 60}], "warnings": []}
        texts = [c.value for row in load_workbook(io.BytesIO(build_report_excel(rep, "t")))["Báo cáo"].iter_rows()
                 for c in row if isinstance(c.value, str)]
        self.assertIn("Phiên từ AI", texts)
        self.assertIn("Theo nguồn: ChatGPT 60", texts)
        rep.ai = {"error": "Tài khoản chưa cấp quyền"}
        texts = [c.value for row in load_workbook(io.BytesIO(build_report_excel(rep, "t")))["Báo cáo"].iter_rows()
                 for c in row if isinstance(c.value, str)]
        self.assertTrue(any("Chưa lấy được số liệu GA4" in t for t in texts))


if __name__ == "__main__":
    unittest.main()


class MonthlyReportTest(unittest.TestCase):
    def setUp(self):
        self.con = db.connect(":memory:")
        # Tháng 9 và tháng 10 trọn; dữ liệu đến 02/11 nên tháng 10 đã đủ.
        for d, a_clicks, b_clicks in [("2026-09-10", 40, 30), ("2026-10-10", 70, 10), ("2026-10-31", 10, 0)]:
            db.replace_day(self.con, SITE, d, [page(A, a_clicks), page(B, b_clicks)],
                           [{"keys": [A, "gọng kính"], "clicks": a_clicks, "impressions": 100, "ctr": 0,
                             "position": 2.0}])
        db.replace_day(self.con, SITE, "2026-11-02", [], [])
        db.save_site_totals(self.con, SITE, [
            {"keys": [d], "clicks": c, "impressions": c * 8, "ctr": 0, "position": 5.0}
            for d, c in [("2026-09-10", 75), ("2026-10-10", 85), ("2026-10-31", 12)]])
        U.assign_urls(self.con, SITE)
        self.gid = U.ensure_group(self.con, SITE)
        U.add_to_group(self.con, SITE, self.gid, f"{A}\n{B}", "test")
        self.cfg = Config(site_url=SITE, db_path=Path(":memory:"), client_secret=Path("x"), token_path=Path("y"))
        R.add_keywords(self.con, SITE, self.cfg.rank, report.date(2026, 11, 2), [("gọng kính", A)], "test")
        kid = R.keywords(self.con, SITE, report.date(2026, 11, 2))[0]["id"]
        for d, rank in [("2026-09-08", 9), ("2026-09-29", 6), ("2026-10-06", 4), ("2026-10-27", 3)]:
            self.con.execute("INSERT INTO rank_checks VALUES (?, ?, 30, ?, ?, NULL, 0.0018, 'x')", (kid, d, rank, A))
        self.rep = report.build_month(self.con, self.cfg, self.gid, 0, with_ai=False)

    def test_months_only_complete(self):
        self.assertEqual(report.months(report.date(2026, 11, 2), 2)[0],
                         (report.date(2026, 10, 1), report.date(2026, 10, 31)))
        self.assertEqual(report.months(report.date(2026, 10, 31), 1)[0][0], report.date(2026, 10, 1))
        self.assertEqual(report.months(report.date(2026, 10, 30), 1)[0][0], report.date(2026, 9, 1))  # T10 chưa trọn

    def test_month_kpis_from_daily_not_weeks(self):
        r = self.rep
        self.assertEqual((r.week, r.prev_week), ((report.date(2026, 10, 1), report.date(2026, 10, 31)),
                                                 (report.date(2026, 9, 1), report.date(2026, 9, 30))))
        self.assertEqual((r.site_kpi.clicks, r.site_prev.clicks), (97, 75))  # tổng property, cả ngày 31/10
        self.assertEqual((r.group_kpi.clicks, r.group_prev.clicks), (90, 70))
        self.assertEqual(r.trend[-1]["week"], "T10/26")

    def test_month_rank_last_check_vs_last_of_prev_month(self):
        r = self.rep.ranks[0]
        self.assertEqual((r["rank"], r["prev"], r["date"], r["prev_date"]), (3, 6, "2026-10-27", "2026-09-29"))

    def test_month_extras_and_excel(self):
        r = self.rep
        self.assertEqual([(x["url"], x["clicks"], x["prev"]) for x in r.top_urls], [(A, 80, 40), (B, 10, 30)])
        self.assertEqual(r.top_queries[0]["query"], "gọng kính")
        self.assertEqual(r.top_queries[0]["clicks"], 80)
        report.save_note(self.con, SITE, report.note_key(r), "Tháng tốt.", "test")
        report.save_note(self.con, SITE, "plan:2026-10-01", "Viết 5 bài.", "test")
        rep = report.build_month(self.con, self.cfg, self.gid, 0, with_ai=False)
        self.assertEqual((rep.note, rep.plan), ("Tháng tốt.", "Viết 5 bài."))
        self.assertEqual(report.get_note(self.con, SITE, report.date(2026, 10, 1)), "")  # không lẫn với tuần
        ws = load_workbook(io.BytesIO(build_report_excel(rep, "test")))["Báo cáo"]
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        for t in ("Báo cáo SEO & GEO tháng 10/2026", "Kế hoạch tháng sau", "Viết 5 bài.", "Theo loại trang (cộng theo URL)",
                  "Top 10 từ khóa (query) nhiều click nhất tháng", "Click tháng này"):
            self.assertIn(t, texts)


class RankChangeTextTest(unittest.TestCase):
    def test_change_text(self):
        from seo_report.web_rank import change_text
        base = {"date": "2026-10-27", "rank": 3, "prev": 5, "prev_depth": 30, "last_depth": 10}
        self.assertEqual(change_text(base)[0], "▲ 2")
        self.assertEqual(change_text({**base, "rank": 7})[0], "▼ 2")
        self.assertEqual(change_text({**base, "prev": None, "prev_depth": 10})[0], "▲ vào top 10")
        self.assertEqual(change_text({**base, "rank": None})[0], "▼ rớt khỏi top 10")
        self.assertEqual(change_text({**base, "prev_depth": None, "prev": None})[0], "Mới")
        self.assertEqual(change_text({**base, "date": None})[0], "")


class SummaryTest(unittest.TestCase):
    def test_seo_geo_summary(self):
        t = ReportTest()
        t.setUp()
        sm = report.summary(t.rep)
        seo = {x["short"]: x for x in sm["seo"]}
        self.assertEqual((seo["Click"]["value"], seo["Click"]["delta"]), ("160", "+23,1%"))
        self.assertEqual(seo["Top 10"]["value"], "1/1")
        self.assertEqual(seo["Top 10"]["delta"], "")  # hạng 5 → 2: vẫn trong top 10, không đổi số
        geo = {x["short"]: x for x in sm["geo"]}
        self.assertEqual(geo["Phiên AI"]["value"], "chưa có")  # chưa kết nối GA4: không ghi 0
        self.assertEqual(geo["AI Overviews"]["value"], "chưa có")
        notes = [n["long"] for n in sm["seo_notes"]]
        self.assertIn("URL tăng click nhiều nhất: /gong-kinh/ (+20 click)", notes)
        self.assertIn("1 keyword đang xếp hạng bằng URL khác URL mục tiêu", notes)
        self.assertEqual(sm["seo_notes"][0]["short"], "/gong-kinh/ +20")
