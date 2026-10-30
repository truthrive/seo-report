"""Nhập file export báo cáo Generative AI của GSC (AI Overviews / AI Mode). GSC không có API cho báo cáo này, chỉ có nút
Export (đã kiểm tra tài liệu Google ngày 30/10/2026).

File export (zip CSV hoặc Excel) gồm 5 bảng: Chart (ngày → impression), Pages (≤ 1.000 trang), Countries, Devices,
Filters. Bảng được nhận ra theo NỘI DUNG chứ không theo tên file/tiêu đề cột, để chạy được cả khi GSC để tiếng Việt.

Quy tắc dữ liệu (spec §4.3, tài liệu Google):
- Ô "~" / "-" trong GSC bị xuất thành 0, nên mọi số 0 lưu là NULL (không rõ), không cộng như số 0 thật.
- Chart, Countries, Devices tính theo property; Pages tính theo trang. Không cộng chéo giữa các bảng.
- File không ghi ngày cụ thể ("Last 3 months"): kỳ lấy từ ngày đầu/cuối của bảng Chart.
- Nhập trùng ngày: số theo ngày lấy từ lần nhập mới nhất (số gần đây của GSC còn thay đổi). Bảng trang/quốc gia/thiết bị
  thuộc riêng kỳ của từng lần nhập, không gộp giữa các lần nhập.
- Xóa một lần nhập là ẩn đi (giữ lại để đối soát), số theo ngày tự lấy lại từ lần nhập cũ hơn nếu có.
"""
import csv
import io
import json
import re
import sqlite3
import zipfile
from datetime import date, datetime, timezone
from urllib.parse import urlsplit

from . import urls as U

SCHEMA = """
CREATE TABLE IF NOT EXISTS gsc_ai_imports (
    id INTEGER PRIMARY KEY,
    site TEXT NOT NULL,
    imported_at TEXT NOT NULL,
    user TEXT NOT NULL,
    file_name TEXT NOT NULL,
    search_type TEXT NOT NULL,
    period_start TEXT NOT NULL,
    period_end TEXT NOT NULL,
    filters TEXT NOT NULL,
    stats TEXT NOT NULL,
    deleted_at TEXT
);
CREATE TABLE IF NOT EXISTS gsc_ai_daily (
    import_id INTEGER NOT NULL,
    date TEXT NOT NULL,
    impressions INTEGER,
    PRIMARY KEY (import_id, date)
) WITHOUT ROWID;
CREATE TABLE IF NOT EXISTS gsc_ai_breakdown (
    import_id INTEGER NOT NULL,
    scope TEXT NOT NULL,
    key TEXT NOT NULL,
    url_id INTEGER,
    impressions INTEGER,
    PRIMARY KEY (import_id, scope, key)
) WITHOUT ROWID;
"""
PAGE_ROW_LIMIT = 1000
DEVICE_WORDS = {"mobile", "desktop", "tablet", "di động", "máy tính", "máy tính bảng", "điện thoại"}
STANDARD_FILTER_WORDS = ("search type", "loại tìm kiếm", "date", "ngày")
MULTIMODAL_WORDS = ("multimodal", "đa phương thức")


class ImportError_(ValueError):
    """File không đúng định dạng / không đúng property: báo lại người dùng, không nhập gì."""


def init(con: sqlite3.Connection) -> None:
    con.executescript(SCHEMA)


# ---- Đọc file ----

def _tables_from_upload(name: str, content: bytes) -> list[list[list[str]]]:
    if name.lower().endswith(".zip"):
        tables = []
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            for info in z.infolist():
                if info.filename.lower().endswith(".csv") and info.file_size < 20_000_000:
                    text = z.read(info).decode("utf-8-sig")
                    tables.append([r for r in csv.reader(io.StringIO(text)) if r])
        return tables
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [[[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)
                 if any(v is not None for v in row)] for ws in wb.worksheets]
    raise ImportError_("Chỉ nhận file .zip (Export → Download CSV) hoặc .xlsx (Export → Download Excel) từ GSC.")


def _cell(v) -> str:
    """Ô Excel → chữ. Số Excel đọc ra dạng float (4630.0): đổi về số nguyên trước, nếu không bước bỏ dấu phân cách
    nghìn ở _num sẽ biến 4630.0 thành 46300. Ngày Excel đổi về YYYY-MM-DD."""
    from datetime import date, datetime
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _num(v: str) -> int | None:
    v = (v or "").strip().replace(",", "").replace(".", "")
    if not re.fullmatch(r"\d+", v):
        raise ValueError(v)
    return int(v)


def _kind(rows: list[list[str]]) -> str:
    """Nhận ra bảng theo nội dung cột đầu (không phụ thuộc ngôn ngữ giao diện GSC)."""
    body = [r for r in rows[1:] if r and r[0].strip()]
    if len(rows[0]) == 2 and body and all(len(r) >= 2 for r in body):
        firsts = [r[0].strip() for r in body]
        try:
            [_num(r[1]) for r in body]
        except ValueError:
            return "filters"
        if all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", f) for f in firsts):
            return "chart"
        if all(f.lower().startswith(("http://", "https://")) for f in firsts):
            return "pages"
        if all(f.lower() in DEVICE_WORDS for f in firsts):
            return "devices"
        return "countries"
    if len(rows) == 1:
        return "empty"
    return "unknown"


def parse(name: str, content: bytes, site: str) -> dict:
    tables = _tables_from_upload(name, content)
    found: dict[str, list] = {}
    for t in tables:
        if not t:
            continue
        k = _kind(t)
        if k in found:
            raise ImportError_(f"File có hai bảng giống nhau ({k}): không phải file export chuẩn của GSC.")
        found[k] = t
    if "chart" not in found:
        raise ImportError_("Không thấy bảng theo ngày (Chart) trong file. Hãy dùng nút Export trong báo cáo "
                           "Generative AI của GSC, không sửa file trước khi nhập.")
    filters = {r[0].strip(): r[1].strip() for r in found.get("filters", [])[1:] if len(r) >= 2}
    extra = [k for k in filters if not k.lower().startswith(STANDARD_FILTER_WORDS)]
    if extra:
        raise ImportError_("File được xuất khi đang lọc thêm (" + ", ".join(f"{k}: {filters[k]}" for k in extra) +
                           "). Hãy bỏ mọi bộ lọc trừ kỳ và loại tìm kiếm rồi xuất lại, để số là của toàn site.")
    st = next((v for k, v in filters.items() if k.lower().startswith(("search type", "loại tìm kiếm"))), "")
    search_type = "multimodal" if any(w in st.lower() for w in MULTIMODAL_WORDS) else "text"

    def pairs(kind):
        out, unknown = {}, 0
        for r in found.get(kind, [])[1:]:
            if not r or not r[0].strip():
                continue
            v = _num(r[1])
            if v == 0:  # "~" hoặc "-" trong GSC: không rõ, không phải 0
                v, unknown = None, unknown + 1
            if r[0].strip() in out:
                raise ImportError_(f"Bảng {kind} có dòng trùng: {r[0]}")
            out[r[0].strip()] = v
        return out, unknown

    daily, unk_d = pairs("chart")
    pages, unk_p = pairs("pages")
    countries, unk_c = pairs("countries")
    devices, unk_v = pairs("devices")
    host = urlsplit(site).netloc.removeprefix("www.")
    wrong = [p for p in pages if urlsplit(p).netloc.removeprefix("www.") != host]
    if wrong:
        raise ImportError_(f"File có trang không thuộc {host} (ví dụ {wrong[0]}): có thể xuất nhầm property.")
    days = sorted(daily)
    return {"search_type": search_type, "search_type_raw": st, "filters": filters,
            "period_start": days[0], "period_end": days[-1], "daily": daily, "pages": pages,
            "countries": countries, "devices": devices,
            "stats": {"days": len(daily), "pages": len(pages), "countries": len(countries), "devices": len(devices),
                      "unknown_zero": {"chart": unk_d, "pages": unk_p, "countries": unk_c, "devices": unk_v},
                      "pages_capped": len(pages) >= PAGE_ROW_LIMIT,
                      "min_page_impressions": min((v for v in pages.values() if v), default=None),
                      "total_chart": sum(v for v in daily.values() if v),
                      "missing_tables": [k for k in ("pages", "countries", "devices") if k not in found]}}


# ---- Lưu, xóa ----

def save(con: sqlite3.Connection, site: str, parsed: dict, file_name: str, user: str) -> int:
    init(con)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with con:
        cur = con.execute(
            "INSERT INTO gsc_ai_imports (site, imported_at, user, file_name, search_type, period_start, period_end, "
            "filters, stats) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (site, now, user, file_name, parsed["search_type"], parsed["period_start"], parsed["period_end"],
             json.dumps(parsed["filters"], ensure_ascii=False), json.dumps(parsed["stats"], ensure_ascii=False)))
        iid = cur.lastrowid
        con.executemany("INSERT INTO gsc_ai_daily VALUES (?, ?, ?)", [(iid, d, v) for d, v in parsed["daily"].items()])
        # URL gốc giữ nguyên trong `key`; url_id là URL báo cáo của tool (gộp ?utm, #mục) để ghép với GSC Web.
        rows = [(iid, "page", raw, U._match_existing(con, site, U.normalize(raw)), v)
                for raw, v in parsed["pages"].items()]
        rows += [(iid, "country", k, None, v) for k, v in parsed["countries"].items()]
        rows += [(iid, "device", k, None, v) for k, v in parsed["devices"].items()]
        con.executemany("INSERT INTO gsc_ai_breakdown VALUES (?, ?, ?, ?, ?)", rows)
        U.log_change(con, user, site, "Nhập file GSC Generative AI", None,
                     f"#{iid} {file_name} {parsed['period_start']}..{parsed['period_end']} ({parsed['search_type']})")
    return iid


def delete(con: sqlite3.Connection, site: str, import_id: int, user: str) -> None:
    init(con)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with con:
        con.execute("UPDATE gsc_ai_imports SET deleted_at = ? WHERE id = ? AND site = ? AND deleted_at IS NULL",
                    (now, import_id, site))
        U.log_change(con, user, site, "Xóa lần nhập GSC Generative AI", None, f"#{import_id}")


# ---- Đọc số ----

def imports(con: sqlite3.Connection, site: str, include_deleted: bool = False) -> list[dict]:
    init(con)
    q = "SELECT id, imported_at, user, file_name, search_type, period_start, period_end, stats, deleted_at " \
        "FROM gsc_ai_imports WHERE site = ?" + ("" if include_deleted else " AND deleted_at IS NULL") + \
        " ORDER BY id DESC"
    return [{"id": r[0], "imported_at": r[1], "user": r[2], "file_name": r[3], "search_type": r[4],
             "period_start": r[5], "period_end": r[6], "stats": json.loads(r[7]), "deleted_at": r[8]}
            for r in con.execute(q, (site,))]


def daily(con: sqlite3.Connection, site: str, start: date, end: date, search_type: str = "text") -> dict:
    """Impression AI theo ngày trong kỳ, lấy từ lần nhập mới nhất có ngày đó. Trả về {date: impressions|None} chỉ cho
    các ngày có dữ liệu; ngày không có trong lần nhập nào thì vắng mặt (khác với None = GSC hiện ~)."""
    init(con)
    rows = con.execute(
        """SELECT d.date, d.impressions FROM gsc_ai_daily d JOIN gsc_ai_imports i ON i.id = d.import_id
           WHERE i.site = ? AND i.search_type = ? AND i.deleted_at IS NULL AND d.date BETWEEN ? AND ?
             AND d.import_id = (SELECT MAX(d2.import_id) FROM gsc_ai_daily d2 JOIN gsc_ai_imports i2
                                ON i2.id = d2.import_id WHERE d2.date = d.date AND i2.site = i.site
                                AND i2.search_type = i.search_type AND i2.deleted_at IS NULL)""",
        (site, search_type, start.isoformat(), end.isoformat())).fetchall()
    return dict(rows)


def period_total(con: sqlite3.Connection, site: str, start: date, end: date, search_type: str = "text") -> dict:
    """Tổng impression AI của kỳ + độ phủ. status: NOT_CONNECTED (chưa nhập file nào) / PARTIAL (thiếu ngày hoặc có
    ô không rõ) / MANUAL_IMPORT (đủ). Thiếu ngày thì không trả tổng, tránh hiện số nhỏ hơn thực tế."""
    if not imports(con, site):
        return {"status": "NOT_CONNECTED", "impressions": None, "days": 0, "covered": 0, "unknown": 0}
    d = daily(con, site, start, end, search_type)
    n = (end - start).days + 1
    unknown = sum(1 for v in d.values() if v is None)
    total = sum(v for v in d.values() if v)
    full = len(d) == n and not unknown
    return {"status": "MANUAL_IMPORT" if full else "PARTIAL", "impressions": total if len(d) == n else None,
            "known_sum": total, "days": n, "covered": len(d), "unknown": unknown}


def breakdown(con: sqlite3.Connection, import_id: int, scope: str) -> list[dict]:
    init(con)
    return [{"key": k, "url_id": u, "impressions": v} for k, u, v in con.execute(
        "SELECT key, url_id, impressions FROM gsc_ai_breakdown WHERE import_id = ? AND scope = ? "
        "ORDER BY impressions IS NULL, impressions DESC", (import_id, scope))]


def web_daily(con: sqlite3.Connection, site: str, start: date, end: date) -> dict:
    """Impression GSC Web toàn site theo ngày (tổng property) để tính tỷ lệ AI Overviews/AI Mode trên Web."""
    return dict(con.execute("SELECT date, impressions FROM gsc_site_daily WHERE site = ? AND date BETWEEN ? AND ?",
                            (site, start.isoformat(), end.isoformat())).fetchall())
