"""Bộ slide báo cáo tuần / tháng (16:9) từ cùng dữ liệu với màn Báo cáo và file Excel. Tải lên Drive thì Google tự
chuyển sang Google Slides; PDF xuất từ chính bộ Slides đó.

Slide: bìa · kết luận (+ kế hoạch tháng sau, điểm cần chú ý) · KPI cam kết · kết quả SEO · kết quả GEO · thứ hạng từ
khóa cam kết · công việc · (tháng) phụ lục: thứ hạng từng từ khóa, theo loại trang, top từ khóa, top trang.
"""
import io
from datetime import datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import ai_format as F
from .projects import domain
from .report import WeeklyReport, change, rank_change, rank_summary, summary

GREEN, LIME, INK, MUTED, RED, CREAM = "1E7C31", "B8D332", "2A2A24", "6B6A5E", "C2410C", "F4F3EE"
AMBER, BLUE = "8F5F14", "2563EB"
FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)
ASSETS = __import__("pathlib").Path(__file__).parent / "assets"
LOGO_WIDE, LOGO_MARK = ASSETS / "logo-wide.png", ASSETS / "logo-mark.png"


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = _rgb(color)
    return box


def _slide(prs, title: str, subtitle: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(GREEN)
    bar.line.fill.background()
    _text(s, Inches(0.6), Inches(0.35), Inches(11), Inches(0.7), title, 26, True)
    if LOGO_MARK.exists():
        s.shapes.add_picture(str(LOGO_MARK), W - Inches(1.15), Inches(0.35), height=Inches(0.62))
    if subtitle:
        _text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.4), subtitle, 12, color=MUTED)
    return s


def _kpi(slide, x, y, label, value, delta="", good=None, w=Inches(2.9), h=Inches(1.45)):
    card = slide.shapes.add_shape(5, x, y, w, h)  # hình chữ nhật bo góc
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb("FFFFFF")
    card.line.color.rgb = _rgb("E7E5DC")
    card.adjustments[0] = 0.08
    _text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.3), Inches(0.35), label, 11, color=MUTED)
    _text(slide, x + Inches(0.2), y + Inches(0.45), w - Inches(0.3), Inches(0.6), value, 26, True)
    if delta:
        _text(slide, x + Inches(0.2), y + Inches(1.0), w - Inches(0.3), Inches(0.35), delta, 12, True,
              MUTED if good is None else (GREEN if good else RED))


def _table(slide, x, y, w, headers, rows, widths=None, font=11, colors=None):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, Inches(0.32) * (len(rows) + 1))
    t = shape.table
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = _rgb(CREAM)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = "" if v is None else str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb("FFFFFF")
    for i in range(len(rows) + 1):
        for j in range(len(headers)):
            for p in t.cell(i, j).text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 or (widths and widths[j] > 3) else PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size, r.font.name = Pt(font), FONT
                    r.font.bold = i == 0
                    color = (colors[i - 1][j] if colors and i > 0 and colors[i - 1][j] else
                             (MUTED if i == 0 else INK))
                    r.font.color.rgb = _rgb(color)
    if widths:
        total = sum(widths)
        for j, wj in enumerate(widths):
            t.columns[j].width = Emu(int(w * wj / total))
    return shape


def _pct_delta(now, prev) -> tuple[str, bool | None]:
    d = change(now, prev)
    return F.delta_pct(d), (None if not d else d > 0)


def _good(ch, prev) -> tuple[str, bool | None]:
    """Nhãn % thay đổi (kỳ trước quá nhỏ thì ghi số kỳ trước) và tốt/xấu."""
    return F.delta_pct(ch, prev), (None if not ch else ch > 0)


def _vn(iso) -> str:
    return f"{iso[8:10]}/{iso[5:7]}" if iso else "–"


def _kpi_status(ch: dict, running: bool) -> tuple[str, bool | None]:
    if ch["target"] is None or ch["ok"] is None:
        return ("không cam kết" if ch["target"] is None else "chưa có số"), None
    if ch["ok"]:
        return ("đang đạt" if running else "đạt"), True
    gap = ch["gap"]
    return (f"thiếu {gap} từ khóa" if "need" in ch else f"thiếu {F.n(gap)} phiên"), False


def _short(url, site):
    return (url or "").replace(site.rstrip("/"), "") or ("/" if url else "")


def build_report_pptx(rep: WeeklyReport, user: str) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    month = rep.kind == "month"
    unit = "tháng" if month else "tuần"
    (ws, we), (ps, pe) = rep.week, rep.prev_week
    site = domain(rep.site)
    period = f"tháng {ws:%m/%Y}" if month else f"tuần {ws:%d/%m} – {we:%d/%m/%Y}"
    prev_text = f"tháng {ps:%m/%Y}" if month else f"tuần {ps:%d/%m} – {pe:%d/%m}"

    # ---- Bìa ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if LOGO_WIDE.exists():
        s.shapes.add_picture(str(LOGO_WIDE), Inches(0.9), Inches(0.8), height=Inches(1.25))
    band = s.shapes.add_shape(1, 0, Inches(5.9), W, Inches(1.6))
    band.fill.solid()
    band.fill.fore_color.rgb = _rgb(GREEN)
    band.line.fill.background()
    accent = s.shapes.add_shape(1, Inches(0.9), Inches(2.75), Inches(1.2), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(LIME)
    accent.line.fill.background()
    _text(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.0), f"Báo cáo SEO & GEO {period}", 38, True)
    _text(s, Inches(0.9), Inches(4.05), Inches(11.5), Inches(0.6), f"{site} · so với {prev_text}",
          18, color=MUTED)
    _text(s, Inches(0.9), Inches(6.45), Inches(11.5), Inches(0.5),
          f"Xuất lúc {datetime.now():%d/%m/%Y %H:%M} bởi {user}", 12, color="E7F3EA")

    # Thứ tự giống màn Báo cáo: Kết luận → KPI cam kết → SEO → GEO → Thứ hạng → Công việc → (tháng) Phụ lục.
    x0, gap, cw = Inches(0.6), Inches(4.1), Inches(3.95)
    sm = summary(rep)
    notes = sm["seo_notes"][:3]

    # ---- Kết luận (chưa viết nhận xét thì thành "Điểm chính": không để chữ "chưa có" trong bản gửi đi) ----
    plan = rep.plan if month else ""
    if rep.note or plan or notes:
        s = _slide(prs, ("Kết luận" if rep.note else "Điểm chính") + (" và kế hoạch tháng sau" if plan else ""),
                   f"{period} · so với {prev_text}")
        y = Inches(1.5)
        if rep.note or plan:
            if rep.note:
                _text(s, x0, y, Inches(6.0) if plan else Inches(12.1), Inches(3.3), rep.note, 16)
            if plan:
                px = Inches(6.9) if rep.note else x0
                _text(s, px, y, Inches(5.8), Inches(0.45), "Kế hoạch tháng sau", 15, True, GREEN)
                _text(s, px, y + Inches(0.45), Inches(5.8), Inches(2.9), plan, 14)
            y = Inches(5.15)
        if notes:
            _text(s, x0, y, Inches(12), Inches(0.4), "Điểm cần chú ý", 14, True)
            for j, note in enumerate(notes):
                mark = {"up": "▲", "down": "▼", "warn": "⚠"}.get(note["tone"], "•")
                color = {"up": GREEN, "down": RED, "warn": AMBER}.get(note["tone"], MUTED)
                _text(s, x0, y + Inches(0.45) + Inches(0.45) * j, Inches(12), Inches(0.4), f"{mark}  {note['long']}",
                      14, color=color)

    # ---- KPI cam kết (khi dự án đã bắt đầu) ----
    k = rep.kpi
    if k and k.get("state") not in ("before", "no_start", None) and any(m["status"] != "future" for m in k["months"]):
        m = next((x for x in k["months"] if x["status"] == "running"),
                 [x for x in k["months"] if x["status"] == "closed"][-1])
        running = m["status"] == "running"
        s = _slide(prs, f"KPI cam kết · Tháng {m['n']}/{len(k['months'])}",
                   f"{_vn(m['from'])} – {_vn(m['to'])} · {m['rank']['total']} từ khóa cam kết · "
                   + ("đang chạy, số còn thay đổi" if running else "đã chốt"))
        for i, key in enumerate(("top3", "top5", "top10", "sessions")):
            ch = m["checks"][key]
            if key == "sessions":
                label, value = "Phiên Organic (GA4)", F.n(ch["actual"]) if ch["actual"] is not None else "chưa có"
                target = F.n(ch["target"]) if ch["target"] is not None else "không cam kết"
            else:
                label = f"Từ khóa top {key[3:]}"
                value = f"{(ch['actual'] or 0) * 100:.0f}%" if ch["actual"] is not None else "–"
                target = f"{ch['target'] * 100:.0f}%" if ch["target"] is not None else "không cam kết"
            status, good = _kpi_status(ch, running)
            _kpi(s, x0 + Inches(3.08) * i, Inches(1.5), label, value, f"Cam kết {target} · {status}", good,
                 w=Inches(2.95), h=Inches(1.5))
        heads = ["Chỉ số"] + [f"T{x['n']}" for x in k["months"]]
        rows, colors = [], []
        for key, lab in (("top3", "Top 3"), ("top5", "Top 5"), ("top10", "Top 10"), ("sessions", "Phiên Organic")):
            tgt, act, col = [f"{lab} · cam kết"], ["Thực tế"], [None]
            for x in k["months"]:
                t = x.get(key)
                tgt.append("–" if t is None else (F.n(t) if key == "sessions" else f"{t * 100:.0f}%"))
                ch = (x.get("checks") or {}).get(key)
                if not ch or x["status"] == "future":
                    act.append("")
                    col.append(None)
                    continue
                act.append(F.n(ch["actual"]) if key == "sessions" and ch["actual"] is not None else
                           (f"{ch['actual'] * 100:.0f}%" if key != "sessions" and ch["actual"] is not None else "–"))
                col.append(None if ch["ok"] is None else (GREEN if ch["ok"] else (AMBER if x["status"] == "running"
                                                                                  else RED)))
            rows += [tgt, act]
            colors += [[None] * len(tgt), col]
        _table(s, x0, Inches(3.3), Inches(12.1), heads, rows, [2.6] + [1.0] * len(k["months"]), 10, colors)

    # ---- Kết quả SEO ----
    s = _slide(prs, "Kết quả SEO", f"Traffic từ Google Search · {period} · so với {prev_text}")
    o = rep.organic if rep.organic and not rep.organic.get("error") else None
    sk, sp = rep.site_kpi, rep.site_prev
    if o:
        _kpi(s, x0, Inches(1.5), "Phiên Organic Search (GA4)", F.n(o["sessions"]), *_good(o["change"], o["prev"]),
             w=cw)
    else:
        _kpi(s, x0, Inches(1.5), "Phiên Organic Search (GA4)", "chưa có", "Chưa kết nối GA4", None, w=cw)
    _kpi(s, x0 + gap, Inches(1.5), "Click (GSC)", F.n(sk.clicks), *_pct_delta(sk.clicks, sp.clicks), w=cw)
    _kpi(s, x0 + gap * 2, Inches(1.5), "Impression (GSC)", F.n(sk.impressions),
         *_pct_delta(sk.impressions, sp.impressions), w=cw)
    data = CategoryChartData()
    data.categories = [t["week"] for t in rep.trend]
    if o:
        data.add_series("Phiên Organic (GA4)", o["spark"])
    data.add_series("Click (GSC)", [t["site_clicks"] for t in rep.trend])
    chart = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x0, Inches(3.2), Inches(12.1), Inches(3.55), data).chart
    chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
    for series, color in zip(chart.series, ([GREEN, BLUE] if o else [BLUE])):
        series.format.line.color.rgb = _rgb(color)
        series.format.line.width = Pt(2.25)
        series.smooth = True
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = _rgb(color)
    chart.font.size, chart.font.name = Pt(11), FONT
    _text(s, x0, Inches(6.8), Inches(12.1), Inches(0.35), "Click đếm lượt bấm từ Google (GSC); phiên đếm lượt vào site "
          "(GA4). Hai nguồn đo khác nhau nên số không bằng nhau.", 10, color=MUTED)

    # ---- Kết quả GEO ----
    s = _slide(prs, "Kết quả GEO", "Hiện diện trên AI Search · hai nguồn đo riêng, không cộng với nhau")
    ai, aio = rep.ai, rep.aio
    if ai and not ai.get("error"):
        kk, dl, pv = ai["kpi"], ai["deltas"], ai["prev"]
        top = ai["sources"][0] if ai["sources"] else None
        total = sum(x["sessions"] for x in ai["sources"]) or 1
        _kpi(s, x0, Inches(1.5), "Phiên từ AI assistant (GA4)", F.n(kk["sessions"]),
             *_good(dl.get("sessions"), pv.get("sessions")), w=cw)
        if top:
            _text(s, x0, Inches(3.05), cw, Inches(0.4), f"Nhiều nhất: {F.SOURCE_NAMES.get(top['source'], top['source'])}"
                  f" {top['sessions'] / total * 100:.0f}%", 12, color=MUTED)
    else:
        _kpi(s, x0, Inches(1.5), "Phiên từ AI assistant (GA4)", "chưa có", "Chưa kết nối GA4", None, w=cw)
    if aio and aio["impressions"] is not None:
        _kpi(s, x0 + gap, Inches(1.5), "Hiển thị trong AI Overviews", F.n(aio["impressions"]),
             *_good(aio["change"], aio["prev"]), w=cw)
        dp = (aio["share"] - aio["prev_share"]) if aio["share"] is not None and aio.get("prev_share") is not None else None
        _kpi(s, x0 + gap * 2, Inches(1.5), "Tỷ lệ trên hiển thị Google", F.pct(aio["share"], 1), F.delta_points(dp),
             None if not dp else dp > 0, w=cw)
    else:
        why = "Chưa nhập file từ GSC" if not aio else f"File GSC mới có số tới {_vn(str(aio['data_until']))}"
        _kpi(s, x0 + gap, Inches(1.5), "Hiển thị trong AI Overviews", "chưa có", why, None, w=cw)
    spark = (aio or {}).get("spark") or []
    labels = [t["week"] for t in rep.trend][-len(spark):] if spark else []
    first = next((i for i, v in enumerate(spark) if v is not None), None)
    if first is not None:
        data = CategoryChartData()
        data.categories = labels[first:]
        data.add_series("Hiển thị trong AI Overviews", spark[first:])
        chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x0, Inches(3.55), Inches(12.1), Inches(2.85),
                                   data).chart
        chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
        ser = chart.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb("C9A46A")
        chart.font.size, chart.font.name = Pt(11), FONT
    _text(s, x0, Inches(6.5), Inches(12.1), Inches(0.7),
          "Phiên AI: người bấm từ ChatGPT, Gemini, Perplexity… vào site (GA4). AI Overviews: số lần trang hiện trong câu "
          "trả lời AI của Google (GSC, chỉ có lượt hiển thị, không có click).", 10, color=MUTED)

    # ---- Thứ hạng ----
    rs = rank_summary(rep, rep.committed)
    if rs["total"]:
        s = _slide(prs, "Thứ hạng từ khóa cam kết" if rs["committed"] else "Thứ hạng từ khóa theo dõi",
                   f"{rs['total']} từ khóa · DataForSEO, Google mobile, Việt Nam · lần check mới nhất so với lần trước")
        for i, n in enumerate((3, 5, 10)):
            diff = rs["now"][n] - rs["prev"][n] if rs["has_prev"] else None
            _kpi(s, x0 + gap * i, Inches(1.5), f"Top {n}", f"{rs['now'][n]}/{rs['total']}",
                 (f"{diff:+d} từ khóa" if diff else "không đổi") if diff is not None else "lần check đầu",
                 None if not diff else diff > 0, w=cw)
        if not rs["has_prev"]:  # chưa có lần trước để so: hiện vị trí từng từ khóa
            shown = [r for r in rep.ranks if not rep.committed or r["keyword"].lower() in rep.committed][:9]
            def wrong(r):
                return bool(r["target_url"] and r["ranking_url"] and r["ranking_url"] != r["target_url"])
            _table(s, x0, Inches(3.3), Inches(12.1), ["Từ khóa", "Vị trí", "Trang đang xếp hạng"],
                   [[r["keyword"], r["rank"] if r["rank"] is not None else f">{r['depth']}",
                     (_short(r["ranking_url"], rep.site) or "–") + ("   ⚠ khác trang mục tiêu" if wrong(r) else "")]
                    for r in shown], [4.2, 1.2, 6.7], 12, [[None, None, AMBER if wrong(r) else None] for r in shown])
            more = rs["total"] - len(shown)
            _text(s, x0, Inches(6.75), Inches(12), Inches(0.4), "Lần check đầu tiên của bộ từ khóa: từ lần sau sẽ có "
                  "từ khóa tăng / giảm mạnh nhất." + (f" Còn {more} từ khóa: xem phụ lục trong file Excel." if more > 0 else ""),
                  11, color=MUTED)
        else:
            for i, (title, items, color) in enumerate((("Tăng mạnh nhất", rs["up"], GREEN),
                                                        ("Giảm mạnh nhất", rs["down"], RED))):
                x = x0 + Inches(6.2) * i
                _text(s, x, Inches(3.3), Inches(5.8), Inches(0.4), title, 15, True, color)
                lines = [f"{mm['keyword']}:  {mm['from'] if mm['from'] is not None else '>' + str(mm['prev_depth'])}"
                         f" → {mm['to'] if mm['to'] is not None else '>' + str(mm['depth'])}" for mm in items]
                _text(s, x, Inches(3.8), Inches(5.8), Inches(2.5), "\n".join(lines) or "Không có.", 13)

    # ---- Công việc ----
    if rep.okr:
        from . import okr as O
        kk = rep.okr
        p, cnt = kk["period"], kk["counts"]
        pct = cnt[O.DONE] / kk["total"] * 100 if kk["total"] else 0
        s = _slide(prs, "Công việc", f"OKR · hiện có {pct:.0f}% hạng mục đã hoàn thành"
                   + (f" · {len(kk['late'])} trễ hạn" if kk["late"] else ""))
        for i, (title, its, color) in enumerate(((f"Hoàn thành trong {unit}", p["done"], GREEN),
                                                 (f"Kế hoạch {unit} sau", p["next"], INK))):
            x = x0 + Inches(6.2) * i
            _text(s, x, Inches(1.5), Inches(5.8), Inches(0.4), f"{title} · {len(its)}", 15, True, color)
            lines = [f"{it['code']}   {it['title']}" for it in its[:10]]
            if len(its) > 10:
                lines.append(f"và {len(its) - 10} hạng mục khác")
            _text(s, x, Inches(2.0), Inches(5.8), Inches(4.8), "\n".join(lines) or "Không có.", 13)

    # ---- Phụ lục (báo cáo tháng) ----
    if month:
        if rep.ranks:
            s = _slide(prs, "Phụ lục · Thứ hạng từng từ khóa", "DataForSEO, Google mobile, Việt Nam")
            rows, colors = [], []
            for r in rep.ranks[:14]:
                text, color, _ = rank_change(r["date"], r["rank"], r["depth"], r["prev"], r["prev_depth"])
                rows.append([r["keyword"], r["rank"] if r["rank"] is not None else f">{r['depth']}", text,
                             _short(r["ranking_url"], rep.site)])
                colors.append([None, None, color.lstrip("#"), None])
            _table(s, x0, Inches(1.45), Inches(12.1), ["Từ khóa", "Vị trí", "Thay đổi", "Trang đang xếp hạng"],
                   rows, [3.4, 1, 1.6, 6.1], 11, colors)
        by_type = [r for r in rep.by_type if r["clicks"]][:10]
        s = _slide(prs, "Phụ lục · Theo loại trang và top từ khóa", "Click trong tháng · cộng theo trang nên tổng khác "
                                                                    "tổng toàn site")
        _table(s, x0, Inches(1.45), Inches(5.6), ["Loại trang", "Click", "Thay đổi"],
               [[r["type"], F.n(r["clicks"]), F.delta_pct(r["change"])] for r in by_type], [3, 1.2, 1.2],
               10, [[None, None, (GREEN if (r["change"] or 0) > 0 else RED) if r["change"] else None]
                    for r in by_type])
        _table(s, Inches(6.6), Inches(1.45), Inches(6.1), ["Từ khóa (GSC)", "Click", "Thay đổi"],
               [[r["query"], F.n(r["clicks"]), F.delta_pct(r["change"])] for r in rep.top_queries[:10]],
               [3.6, 1.1, 1.2], 10)
        s = _slide(prs, "Phụ lục · Top 10 trang của tháng", "Toàn site, theo click")
        _table(s, x0, Inches(1.45), Inches(12.1), ["Trang", "Loại", "Click", "Thay đổi"],
               [[_short(r["url"], rep.site), r["type"], F.n(r["clicks"]), F.delta_pct(r["change"])]
                for r in rep.top_urls], [7, 2, 1.2, 1.3], 11)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def slides_filename(rep: WeeklyReport) -> str:
    host = domain(rep.site).split(".")[0]
    if rep.kind == "month":
        return f"bao-cao-seo-thang_{host}_{rep.week[0]:%Y-%m}"
    return f"bao-cao-seo_{host}_{rep.week[0]:%Y-%m-%d}_{rep.week[1]:%Y-%m-%d}"
