"""Xuất Excel "dữ liệu" để kiểm tra số liệu: bảng phẳng, không gộp ô, có bộ lọc, làm pivot được.

Các sheet theo spec (mục Xuất Excel).
"""
import io
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import indexing as IX
from . import keywords as K
from . import rank as R
from . import urls as U

HEADER_FILL = PatternFill("solid", fgColor="E4EFDC")
INT, PCT, POS, CHANGE = "#,##0", "0.0%", "0.0", "+0.0%;-0.0%;0.0%"


def _ctr(c, i):
    return c / i if i else None


def _change(now, prev):
    return (now - prev) / prev if prev else None


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, int, str | None]], rows: list[list]):
    """columns: (tiêu đề, độ rộng, định dạng số hoặc None)."""
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for idx, (_, width, fmt) in enumerate(columns, 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_data_export(con, cfg, group_id: int, week_index: int, user: str, today: date | None = None) -> bytes:
    site, s = cfg.site_url, cfg.keywords
    end = K.latest_date(con, site)
    wks = K.weeks(end, 13 + week_index)
    (ws_, we_), (ps, pe) = wks[week_index], wks[week_index + 1]
    group = dict(U.list_groups(con, site))[group_id]
    members = U.group_members(con, group_id)
    ids = [m["id"] for m in members]
    analysis = K.analyze(con, site, ids, end, s)
    targets = K.targets(con, ids)
    now, prev = K.url_period(con, ids, ws_, we_), K.url_period(con, ids, ps, pe)
    win_start = end - timedelta(days=s.window_days - 1)

    wb = Workbook()
    wb.remove(wb.active)

    info = [
        ["Site", site], ["Nhóm URL", f"{group} ({len(ids)} URL)"],
        ["Kỳ", f"{ws_:%d/%m/%Y} – {we_:%d/%m/%Y}"], ["Kỳ so sánh", f"{ps:%d/%m/%Y} – {pe:%d/%m/%Y}"],
        ["Head KW, keyword phụ tính trên", f"{win_start:%d/%m/%Y} – {end:%d/%m/%Y} ({s.window_days} ngày)"],
        ["Ngưỡng keyword phụ", f"impression ≥ {s.min_impressions}, position ≤ {s.max_position:g}, "
                               f"tối đa {s.max_secondary}, loại thương hiệu: {', '.join(s.brand_terms)}"],
        ["Dữ liệu GSC đến ngày", f"{end:%d/%m/%Y}"],
        ["Xuất lúc", f"{datetime.now():%d/%m/%Y %H:%M} bởi {user}"],
        [],
        ["Lưu ý", "GSC trễ 2–3 ngày; số của các ngày gần nhất còn được cập nhật."],
        ["", "Click theo URL gồm cả query Google ẩn, nên luôn lớn hơn tổng click của các keyword (sheet URL_Keyword)."],
        ["", "Position là trung bình có trọng số theo impression, như cách GSC tính."],
        ["", "URL đã gộp #mục và ?utm_…: một dòng URL có thể gồm nhiều dòng trong GSC."],
        ["", "Rank thực tế: DataForSEO, mobile, Việt Nam. Rank là vị trí tốt nhất của site; nếu URL đang xếp hạng "
             "khác URL đã gán, xem cột 'Đúng URL' ở sheet Rank_History_Long."],
    ]
    ws = wb.create_sheet("Thông tin")
    for r in info:
        ws.append(r)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    # Tổng quan theo loại URL, kèm dòng tổng.
    by_type: dict[str, list] = {}
    for m in members:
        c, i, p = now.get(m["id"], (0, 0, None))
        pc, pi, pp = prev.get(m["id"], (0, 0, None))
        t = by_type.setdefault(m["type"] or "Chưa phân loại", [0, 0, 0, 0, 0, 0, 0.0, 0.0, 0])
        t[0] += 1
        t[1] += c; t[2] += pc; t[3] += i; t[4] += pi
        t[6] += (p or 0) * i; t[7] += (pp or 0) * pi
    total = [sum(v[k] for v in by_type.values()) for k in range(8)]
    overview = []
    for name, v in sorted(by_type.items(), key=lambda kv: -kv[1][1]) + [("Tất cả", total)]:
        n, c, pc, i, pi, _, wp, wpp = v[:8]
        overview.append([name, n, c, pc, _change(c, pc), i, pi, _change(i, pi), _ctr(c, i), _ctr(pc, pi),
                         wp / i if i else None, wpp / pi if pi else None])
    _sheet(wb, "Tổng quan", [
        ("Loại URL", 22, None), ("Số URL", 9, INT), ("Click", 10, INT), ("Click kỳ trước", 12, INT),
        ("% click", 10, CHANGE), ("Impression", 12, INT), ("Impression kỳ trước", 14, INT), ("% impression", 12, CHANGE),
        ("CTR", 9, PCT), ("CTR kỳ trước", 11, PCT), ("Position", 10, POS), ("Position kỳ trước", 13, POS),
    ], overview)

    # URL
    url_rows, kw_rows = [], []
    index = IX.latest(con, ids)
    for m in members:
        a = analysis[m["id"]]
        c, i, p = now.get(m["id"], (0, 0, None))
        pc, pi, pp = prev.get(m["id"], (0, 0, None))
        week_q = K.query_period(con, m["id"], ws_, we_)
        prev_q = K.query_period(con, m["id"], ps, pe)
        visible = sum(v[0] for v in week_q.values())
        match = K.head_matches_target(a.head, targets.get(m["id"]))
        url_rows.append([
            m["url"], m["type"], targets.get(m["id"], ""), a.head.query if a.head else "",
            {True: "Khớp", False: "Lệch", None: ""}[match], c, pc, c - pc, _change(c, pc), i, pi, _ctr(c, i),
            _ctr(pc, pi), p, pp, visible, visible / c if c else None, len(a.secondary), len(a.cannibal),
            *_index_cells(index.get(m["id"]), m["url"]),
        ])
        roles = {}
        for k in ([a.head] if a.head else []) + a.secondary + a.excluded:
            roles[k.query_id] = k
        texts = K._texts(con, "queries", "text", set(week_q) - set(roles))
        for qid in list(roles) + sorted(texts):
            k = roles.get(qid)
            wc, wi, wp = week_q.get(qid, (0, 0, None))
            pc_, pi_, pp_ = prev_q.get(qid, (0, 0, None))
            kw_rows.append([
                m["url"], k.query if k else texts[qid], k.role if k else "Ngoài cửa sổ 28 ngày",
                k.reason if k else "", (k.stronger_url or "") if k else "",
                k.clicks if k else None, k.impressions if k else None, k.position if k else None,
                wc, wi, _ctr(wc, wi), wp, pc_, pi_, pp_,
            ])
    _sheet(wb, "URL", [
        ("URL", 60, None), ("Loại", 18, None), ("Target KW", 24, None), ("Head KW (28 ngày)", 28, None),
        ("Head so với target", 12, None), ("Click", 9, INT), ("Click kỳ trước", 12, INT), ("Chênh click", 11, INT),
        ("% click", 9, CHANGE), ("Impression", 12, INT), ("Impression kỳ trước", 14, INT), ("CTR", 8, PCT),
        ("CTR kỳ trước", 11, PCT), ("Position", 9, POS), ("Position kỳ trước", 13, POS),
        ("Click nhìn thấy theo keyword", 14, INT), ("% click nhìn thấy", 12, PCT), ("Số keyword phụ", 10, INT),
        ("Số query cannibalization", 13, INT), ("Trạng thái index", 30, None), ("Rớt index", 9, None),
        ("Crawl gần nhất", 12, "dd/mm/yyyy"), ("Canonical Google chọn (nếu khác URL)", 50, None),
        ("Ngày kiểm tra index", 12, "dd/mm/yyyy"),
    ], url_rows)
    _sheet(wb, "URL_Keyword", [
        ("URL", 55, None), ("Query", 36, None), ("Vai trò", 12, None), ("Lý do loại", 22, None),
        ("URL mạnh hơn", 45, None), ("Click 28 ngày", 11, INT), ("Impression 28 ngày", 13, INT),
        ("Position 28 ngày", 12, POS), ("Click", 9, INT), ("Impression", 11, INT), ("CTR", 8, PCT),
        ("Position", 9, POS), ("Click kỳ trước", 12, INT), ("Impression kỳ trước", 14, INT),
        ("Position kỳ trước", 13, POS),
    ], kw_rows)

    _sheet(wb, "Cannibalization", [
        ("Query", 36, None), ("URL", 60, None), ("URL thuộc nhóm", 12, None), ("URL mạnh nhất", 12, None),
        ("Click 28 ngày", 11, INT), ("Impression 28 ngày", 13, INT), ("Position 28 ngày", 12, POS),
    ], [[r["query"], r["url"], "Có" if r["in_group"] else "Không", "Có" if r["strongest"] else "",
         r["clicks"], r["impressions"], r["position"]] for r in K.cannibal_rows(con, site, ids, end, s)])

    daily_start = we_ - timedelta(days=12 * 7 - 1)
    _rank_sheets(wb, con, ids, daily_start, we_)
    daily = []
    if ids:
        q = f"""SELECT u.url, d.date, SUM(d.clicks), SUM(d.impressions),
                       SUM(d.position * d.impressions) / NULLIF(SUM(d.impressions), 0)
                FROM gsc_page_daily d JOIN page_url pu ON pu.page_id = d.page_id JOIN urls u ON u.id = pu.url_id
                WHERE u.id IN ({','.join('?' * len(ids))}) AND d.date BETWEEN ? AND ?
                GROUP BY u.id, d.date ORDER BY u.url, d.date"""
        for url, d, c, i, p in con.execute(q, (*ids, daily_start.isoformat(), we_.isoformat())):
            daily.append([url, date.fromisoformat(d), c, i, _ctr(c, i), p])
    _sheet(wb, "URL_Daily", [("URL", 60, None), ("Ngày", 12, "dd/mm/yyyy"), ("Click", 9, INT),
                                  ("Impression", 11, INT), ("CTR", 8, PCT), ("Position", 9, POS)], daily)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _index_cells(ix: dict | None, url: str) -> list:
    if not ix:
        return ["Chưa kiểm tra", "", None, "", None]
    crawl = date.fromisoformat(ix["last_crawl"][:10]) if ix["last_crawl"] else None
    other = ix["google_canonical"] if IX.canonical_mismatch(ix, url) else ""
    return [ix["coverage"] or ix["verdict"], "Có" if ix["dropped"] else "", crawl, other,
            date.fromisoformat(ix["checked_at"][:10])]


def _rank_sheets(wb: Workbook, con, ids: list[int], start: date, end: date) -> None:
    """Rank_History: mỗi keyword một dòng, mỗi ngày check một cột (để đọc). Rank_History_Long: mỗi lần check một dòng
    (để pivot)."""
    R.ensure_tables(con)
    checks = []
    if ids:
        checks = con.execute(
            f"""SELECT u.url, k.keyword, c.check_date, c.rank, c.ranking_url, c.target_rank, c.depth, c.cost
                FROM rank_checks c JOIN rank_keywords k ON k.id = c.keyword_id JOIN urls u ON u.id = k.url_id
                WHERE k.url_id IN ({','.join('?' * len(ids))}) AND c.check_date BETWEEN ? AND ?
                ORDER BY u.url, k.keyword, c.check_date""",
            (*ids, start.isoformat(), end.isoformat()),
        ).fetchall()
    dates = sorted({c[2] for c in checks})
    wide: dict[tuple, dict] = {}
    for url, query, d, rank, *_rest in checks:
        depth = _rest[2]
        wide.setdefault((url, query), {})[d] = rank if rank is not None else f">{depth}"
    _sheet(wb, "Rank_History", [("URL", 55, None), ("Keyword", 32, None)] +
           [(date.fromisoformat(d).strftime("%d/%m"), 7, None) for d in dates],
           [[url, query] + [cells.get(d) for d in dates] for (url, query), cells in wide.items()])
    _sheet(wb, "Rank_History_Long", [
        ("URL", 55, None), ("Keyword", 32, None), ("Ngày check", 12, "dd/mm/yyyy"),
        ("Rank (tốt nhất của site)", 12, INT), ("URL đang xếp hạng", 55, None), ("Đúng URL", 10, None),
        ("Rank của URL này", 12, INT), ("Độ sâu", 8, INT), ("Thiết bị", 9, None), ("Địa điểm", 10, None),
        ("Chi phí USD", 10, "0.0000"),
    ], [[url, query, date.fromisoformat(d), rank, ranking_url,
         ("Có" if ranking_url == url else "Không") if ranking_url else "", target_rank, depth, "mobile", "Việt Nam",
         cost] for url, query, d, rank, ranking_url, target_rank, depth, cost in checks])


def data_export_filename(cfg, con, week_index: int) -> str:
    start, end = K.weeks(K.latest_date(con, cfg.site_url), week_index + 1)[week_index]
    host = cfg.site_url.split("//")[-1].strip("/").split(".")[0]
    return f"seo-du-lieu_{host}_{start:%Y-%m-%d}_{end:%Y-%m-%d}.xlsx"
