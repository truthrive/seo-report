"""File Excel báo cáo tuần / tháng gửi quản lý / khách hàng: một sheet định dạng sẵn, đọc từ trên xuống, có biểu đồ xu hướng.

Khác file "Xuất dữ liệu" (bảng phẳng để kiểm tra/pivot): file này để đọc và gửi đi, nên chỉ có số đã chọn lọc.
"""
import io
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .projects import domain
from .report import WeeklyReport, change, rank_change, rank_summary, summary

GREEN = "1E7C31"
RED = "C2410C"
AMBER = "8F5F14"
BLUE = "2563EB"
INK = "2A2A24"
MUTED = "6B6A5E"
TINT = "E4EFDC"
CREAM = "F4F3EE"
THIN = Side(style="thin", color="E6E4DC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _vn(d) -> str:
    return f"{d:%d/%m/%Y}"


def _short(url: str | None, site: str) -> str:
    return (url or "").replace(site.rstrip("/"), "") or ("/" if url else "")


def _pos_text(rank, depth) -> str:
    return str(rank) if rank is not None else (f">{depth}" if depth else "")


class _Sheet:
    """Ghi lần lượt từng khối xuống dưới."""

    def __init__(self, ws):
        self.ws, self.row = ws, 1

    def title(self, text, size=16, color=INK):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(size=size, bold=True, color=color)
        self.row += 1

    def text(self, text, color=MUTED, italic=False):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(color=color, italic=italic, size=10)
        self.row += 1

    def gap(self, n=1):
        self.row += n

    def section(self, text):
        self.gap()
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(size=12, bold=True, color=GREEN)
        self.row += 1

    def table(self, header: list[str], rows: list[list], formats: list[str | None] | None = None,
              colors: list[list[str | None]] | None = None):
        for j, h in enumerate(header, 1):
            c = self.ws.cell(self.row, j, h)
            c.font = Font(bold=True, color=INK)
            c.fill = PatternFill("solid", fgColor=TINT)
            c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
            c.border = BOX
        self.row += 1
        for i, r in enumerate(rows):
            for j, v in enumerate(r, 1):
                c = self.ws.cell(self.row, j, v)
                c.border = BOX
                if formats and formats[j - 1]:
                    c.number_format = formats[j - 1]
                if j > 1:
                    c.alignment = Alignment(horizontal="right" if isinstance(v, (int, float)) else "left")
                color = colors[i][j - 1] if colors else None
                if color:
                    c.font = Font(color=color, bold=True)
            self.row += 1


def _vn_iso(iso) -> str:
    return f"{iso[8:10]}/{iso[5:7]}" if iso else "–"


def _source(code: str) -> str:
    from . import ai_format as F
    return F.SOURCE_NAMES.get(code, code)


def _unit(rep: WeeklyReport) -> str:
    return "tháng" if rep.kind == "month" else "tuần"


def _delta_color(value, lower_is_better=False):
    if value is None or value == 0:
        return None
    good = value < 0 if lower_is_better else value > 0
    return GREEN if good else RED




def build_report_excel(rep: WeeklyReport, user: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo"
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGH", [44, 16, 16, 14, 16, 16, 14, 30]):
        ws.column_dimensions[col].width = width
    s = _Sheet(ws)
    site_name = domain(rep.site)
    (ws_, we), (ps, pe) = rep.week, rep.prev_week

    unit = _unit(rep)
    month = rep.kind == "month"
    logo = Path(__file__).parent / "assets" / "logo-wide.png"
    if logo.exists():
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(logo))
        img.height, img.width = 44, int(44 * img.width / img.height)
        ws.add_image(img, "F1")
    if month:
        s.title(f"Báo cáo SEO & GEO tháng {ws_:%m/%Y}")
        s.text(f"{site_name} · so với tháng {ps:%m/%Y} · xuất lúc {datetime.now():%d/%m/%Y %H:%M} bởi {user}")
    else:
        s.title(f"Báo cáo SEO & GEO tuần {_vn(ws_)} – {_vn(we)}")
        s.text(f"{site_name} · so với tuần {_vn(ps)} – {_vn(pe)} · xuất lúc {datetime.now():%d/%m/%Y %H:%M} bởi {user}")
    if rep.incomplete_days:
        s.text(f"Lưu ý: còn {rep.incomplete_days} ngày trong {unit} chưa có dữ liệu GSC; số có thể tăng thêm.",
               color=RED)

    # Nhận xét lên đầu: người nhận đọc kết luận trước.
    for title, text, empty in [("Nhận xét", rep.note, "(Chưa có nhận xét)")] + (
            [("Kế hoạch tháng sau", rep.plan, "(Chưa có kế hoạch)")] if month else []):
        if not text:  # bản gửi đi không để ô "chưa có"
            continue
        s.section(title)
        ws.merge_cells(start_row=s.row, start_column=1, end_row=s.row + 3, end_column=8)
        c = ws.cell(s.row, 1, text or empty)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor=CREAM)
        c.font = Font(color=INK if text else MUTED, italic=not text)
        s.gap(4)

    # Thứ tự giống màn Báo cáo và slide: Kết luận → KPI cam kết → SEO → GEO → Thứ hạng → Công việc → (tháng) Phụ lục.
    notes = summary(rep)["seo_notes"][:3]
    if notes:
        s.section("Điểm cần chú ý")
        for n in notes:
            s.text({"up": "▲ ", "down": "▼ ", "warn": "⚠ "}.get(n["tone"], "• ") + n["long"],
                   color={"up": GREEN, "down": RED, "warn": AMBER}.get(n["tone"], INK))

    pct_fmt = "+0.0%;-0.0%;0.0%"
    k = rep.kpi
    if k and k.get("state") not in ("before", "no_start", None) and any(m["status"] != "future" for m in k["months"]):
        m = next((x for x in k["months"] if x["status"] == "running"),
                 [x for x in k["months"] if x["status"] == "closed"][-1])
        running = m["status"] == "running"
        s.section(f"KPI cam kết · tháng {m['n']}/{len(k['months'])} ({_vn_iso(m['from'])} – {_vn_iso(m['to'])}"
                  + (", đang chạy" if running else ", đã chốt") + ")")
        rows, colors = [], []
        for key, lab in (("top3", "Từ khóa top 3"), ("top5", "Từ khóa top 5"), ("top10", "Từ khóa top 10"),
                         ("sessions", "Phiên Organic (GA4)")):
            ch = m["checks"][key]
            if key == "sessions":
                tgt, act = ch["target"], ch["actual"]
                extra = f"dự báo cuối tháng {ch['forecast']:,}".replace(",", ".") if running and ch.get("forecast") else ""
            else:
                tgt, act = ch["target"], ch["actual"]
                extra = f"{ch['count']}/{m['rank']['total']} từ khóa"
            status = ("Không cam kết" if tgt is None else "Chưa có số" if ch["ok"] is None else
                      ("Đang đạt" if running else "Đạt") if ch["ok"] else
                      (f"Thiếu {ch['gap']} từ khóa" if key != "sessions" else f"Thiếu {ch['gap']:,} phiên".replace(",", ".")))
            rows.append([lab, tgt, act, status, extra])
            colors.append([None, None, None, None if ch["ok"] is None else (GREEN if ch["ok"] else
                                                                             (AMBER if running else RED)), None])
        s.table(["Chỉ số", "Cam kết", "Thực tế", "Trạng thái", "Ghi chú"], rows, None, colors)
        for i, r in enumerate(rows):  # % cho dòng thứ hạng, số nguyên cho phiên
            fmt = "#,##0" if i == 3 else "0%"
            for col in (2, 3):
                s.ws.cell(s.row - len(rows) + i, col).number_format = fmt

    s.section("Kết quả SEO · traffic từ Google Search")
    o = rep.organic if rep.organic and not rep.organic.get("error") else None
    sk, sp = rep.site_kpi, rep.site_prev
    rows = [["Phiên Organic Search (GA4)", o["sessions"] if o else "chưa có", o["prev"] if o else "",
             o["change"] if o else None],
            ["Click (GSC)", sk.clicks, sp.clicks, change(sk.clicks, sp.clicks)],
            ["Impression (GSC)", sk.impressions, sp.impressions, change(sk.impressions, sp.impressions)]]
    s.table(["Chỉ số", f"{unit.capitalize()} này", f"{unit.capitalize()} trước", "% thay đổi"], rows,
            [None, "#,##0", "#,##0", pct_fmt], [[None, None, None, _delta_color(r[3])] for r in rows])
    s.text("Click đếm lượt bấm từ Google (GSC); phiên đếm lượt vào site (GA4). Hai nguồn đo khác nhau nên số không bằng "
           "nhau.", italic=True)
    s.gap()
    trend_top = s.row
    s.table(["Tháng" if month else "Tuần (bắt đầu)", "Phiên Organic (GA4)", "Click (GSC)"],
            [[t["week"], (o["spark"][i] if o else None), t["site_clicks"]] for i, t in enumerate(rep.trend)],
            [None, "#,##0", "#,##0"])
    chart = LineChart()
    chart.title, chart.height, chart.width = None, 7, 16
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=trend_top, max_row=trend_top + len(rep.trend)),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=trend_top + 1, max_row=trend_top + len(rep.trend)))
    for series, color in zip(chart.series, [GREEN, BLUE]):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 28000
        series.smooth = False
    ws.add_chart(chart, f"E{trend_top}")
    s.gap(max(0, 15 - len(rep.trend) - 1))

    s.section("Kết quả GEO · hiện diện trên AI Search")
    ai, aio = rep.ai if rep.ai and not rep.ai.get("error") else None, rep.aio
    rows = []
    if ai:
        rows.append(["Phiên từ AI assistant (GA4)", ai["kpi"]["sessions"], (ai.get("prev") or {}).get("sessions"),
                     (ai.get("deltas") or {}).get("sessions")])
    else:
        rows.append(["Phiên từ AI assistant (GA4)", "chưa có", "", None])
    if aio and aio["impressions"] is not None:
        rows.append(["Hiển thị trong AI Overviews (GSC)", aio["impressions"], aio["prev"], aio["change"]])
    else:
        rows.append(["Hiển thị trong AI Overviews (GSC)", "chưa có", "", None])
    s.table(["Chỉ số", f"{unit.capitalize()} này", f"{unit.capitalize()} trước", "% thay đổi"], rows,
            [None, "#,##0", "#,##0", pct_fmt], [[None, None, None, _delta_color(r[3])] for r in rows])
    if aio and aio["impressions"] is not None and aio["share"] is not None:
        s.text(f"Tỷ lệ trên tổng hiển thị Google: {aio['share'] * 100:.1f}%".replace(".", ","))
    if ai and ai["sources"]:
        total = sum(x["sessions"] for x in ai["sources"]) or 1
        s.text("Nguồn AI: " + " · ".join(f"{_source(x['source'])} {x['sessions'] / total * 100:.0f}%"
                                         for x in ai["sources"]))
    s.text("Phiên AI: người bấm từ ChatGPT, Gemini… vào site (GA4). AI Overviews: số lần trang hiện trong câu trả lời "
           "AI của Google (GSC, chỉ có lượt hiển thị). Hai nguồn không cộng với nhau.", italic=True)

    rs = rank_summary(rep, rep.committed)
    s.section(("Thứ hạng từ khóa cam kết" if rs["committed"] else "Thứ hạng từ khóa theo dõi")
              + f" · {rs['total']} từ khóa (DataForSEO, Google mobile, Việt Nam)")
    if not rs["total"]:
        s.text(f"Chưa có lần check rank nào trong {unit} này.")
    else:
        rows = [[f"Top {n}", rs["now"][n], rs["prev"][n] if rs["has_prev"] else "",
                 (rs["now"][n] - rs["prev"][n]) if rs["has_prev"] else None] for n in (3, 5, 10)]
        s.table(["Nhóm vị trí", "Lần check này", "Lần check trước", "Chênh lệch"], rows,
                [None, "#,##0", "#,##0", "+#,##0;-#,##0;0"], [[None, None, None, _delta_color(r[3])] for r in rows])
        if not rs["has_prev"]:
            s.text("Đây là lần check đầu tiên của bộ từ khóa: từ lần sau sẽ có từ khóa tăng / giảm mạnh nhất.",
                   italic=True)
        else:
            moves = [["Tăng", m_["keyword"], _pos_text(m_["from"], m_["prev_depth"]), _pos_text(m_["to"], m_["depth"])]
                     for m_ in rs["up"]] + \
                    [["Giảm", m_["keyword"], _pos_text(m_["from"], m_["prev_depth"]), _pos_text(m_["to"], m_["depth"])]
                     for m_ in rs["down"]]
            if moves:
                s.gap()
                s.table(["", "Từ khóa", "Lần trước", "Lần này"], moves, None,
                        [[GREEN if r[0] == "Tăng" else RED, None, None, None] for r in moves])

    if rep.okr:
        from . import okr as O
        kk = rep.okr
        p, cnt = kk["period"], kk["counts"]
        s.section(f"Công việc · OKR hiện có {cnt[O.DONE] / kk['total'] * 100 if kk['total'] else 0:.0f}% hạng mục "
                  "đã hoàn thành"
                  + (f" · {len(kk['late'])} trễ hạn" if kk["late"] else ""))
        rows, colors = [], []
        for title, its, color in ((f"Hoàn thành trong {unit}", p["done"], GREEN), (f"Kế hoạch {unit} sau", p["next"], None)):
            for it in its:
                rows.append([title, it["code"], it["title"]])
                colors.append([color, None, None])
        if rows:
            s.table(["Nhóm", "Mã", "Hạng mục"], rows, None, colors)
        else:
            s.text(f"Không có hạng mục hoàn thành trong {unit} hay kế hoạch {unit} sau.")

    shown = [r for r in rep.ranks if not rep.committed or r["keyword"].lower() in rep.committed]
    if shown:  # tuần và tháng đều có: danh sách đầy đủ cho người muốn xem từng từ khóa
        s.section("Phụ lục · thứ hạng từng từ khóa")
        rows, colors = [], []
        for r in shown:
            text, _, move = rank_change(r["date"], r["rank"], r["depth"], r["prev"], r["prev_depth"])
            rows.append([r["keyword"], f"{_pos_text(r['rank'], r['depth'])} ({r['date'][8:10]}/{r['date'][5:7]})",
                         text, _short(r["ranking_url"], rep.site)])
            colors.append([None, None, _delta_color(move or None), None])
        s.table(["Từ khóa", "Vị trí (ngày check)", "Thay đổi", "Trang đang xếp hạng"], rows, None, colors)
    if month:
        by_type = [r for r in rep.by_type if r["clicks"]]
        s.section("Phụ lục · theo loại trang (cộng theo trang)")
        s.table(["Loại trang", "Click", "Click tháng trước", "% thay đổi", "Impression"],
                [[r["type"], r["clicks"], r["prev_clicks"], r["change"], r["impressions"]] for r in by_type],
                [None, "#,##0", "#,##0", pct_fmt, "#,##0"],
                [[None, None, None, _delta_color(r["change"]), None] for r in by_type])
        s.section("Phụ lục · top 10 trang nhiều click nhất tháng")
        s.table(["Trang", "Loại", "Click", "Click tháng trước", "% thay đổi"],
                [[_short(r["url"], rep.site), r["type"], r["clicks"], r["prev"], r["change"]] for r in rep.top_urls],
                [None, None, "#,##0", "#,##0", pct_fmt],
                [[None, None, None, None, _delta_color(r["change"])] for r in rep.top_urls])
        s.section("Phụ lục · top 10 từ khóa (GSC) nhiều click nhất tháng")
        s.table(["Từ khóa", "Click", "Impression", "Click tháng trước", "% thay đổi"],
                [[r["query"], r["clicks"], r["impressions"], r["prev"], r["change"]] for r in rep.top_queries],
                [None, "#,##0", "#,##0", "#,##0", pct_fmt],
                [[None, None, None, None, _delta_color(r["change"])] for r in rep.top_queries])
        s.text("GSC ẩn các từ khóa hiếm, nên tổng click theo từ khóa nhỏ hơn click toàn site.", italic=True)

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_filename(rep: WeeklyReport) -> str:
    host = domain(rep.site).split(".")[0]
    if rep.kind == "month":
        return f"bao-cao-seo-thang_{host}_{rep.week[0]:%Y-%m}.xlsx"
    return f"bao-cao-seo_{host}_{rep.week[0]:%Y-%m-%d}_{rep.week[1]:%Y-%m-%d}.xlsx"
