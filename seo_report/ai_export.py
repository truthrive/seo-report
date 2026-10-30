"""Xuất Excel màn AI Search & GEO theo đúng bộ lọc đang chọn (spec §3.8: export dùng chính query state, có thông tin
kỳ, nguồn dữ liệu, thời điểm truy vấn, bộ lọc). Dữ liệu là bản đã tải trên màn (cùng cache), không truy vấn khác."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from . import ai_classifier as C
from . import ai_format as F
from .export import INT, _sheet

PCT = "0.00%"
CHG = "+0.0%;-0.0%;0.0%"
SEC = "0.0"


def build_ai_excel(state: dict, cfg, user: str) -> bytes:
    d = state["data"] or {}
    s, src, tr, ke = d.get("summary", {}), d.get("sources", {}), d.get("trend", {}), d.get("key_events")
    land, geo = state.get("land") or {}, state.get("geo") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Thông tin"
    meta = s.get("sourceMeta", {})
    info = [
        ["Kỳ", f"{state['start']:%d/%m/%Y} – {state['end']:%d/%m/%Y}"],
        ["Kỳ so sánh", "{startDate} – {endDate}".format(**s["comparisonPeriod"]) if s.get("comparisonPeriod") else "Không"],
        ["Bộ lọc", f"nguồn = {F.SOURCE_NAMES.get(state['source'], state['source']) or 'tất cả'}; "
                   f"quốc gia = {state['country'] or 'tất cả'}; trang đích = {state.get('landing') or 'tất cả'}"],
        ["Phân loại AI", f"{C.HYBRID} {C.VERSION}: kênh 'AI Assistant' của GA4 + nhận diện bổ sung theo nguồn "
                         "(chatgpt.com, gemini.google.com…); loại traffic có medium trả phí"],
        ["Nguồn dữ liệu", f"GA4 {cfg.ga4_property} (múi giờ của property)"],
        ["Truy vấn lúc (UTC)", meta.get("queriedAt") or ""],
        ["Xuất lúc", f"{datetime.now():%d/%m/%Y %H:%M} bởi {user}"],
        [],
        ["Lưu ý", "Không gồm AI Overviews / AI Mode của Google: GA4 tính phần đó vào Organic Search."],
        ["", "Rate là số GA4 tính (định dạng %), không cộng từ các dòng. Session key event rate = % phiên có ít nhất "
             "1 key event (khác key events / phiên)."],
        ["", "Ô trống = không có số (không có mẫu số hoặc GA4 không tính được), khác với 0."],
        ["", "Cảnh báo GA4: " + (", ".join(sorted({w['code'] for env in (s, src, tr) for w in env.get('warnings', [])}))
                                 or "không có")],
    ]
    for r in info:
        ws.append(r)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 120
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    if s.get("data"):
        a, p, dl = s["data"]["ai"], s["data"].get("aiPrev") or {}, s["data"].get("aiDeltas") or {}
        sh = s["data"]["share"]
        rows = [
            ["Phiên AI", a["sessions"], p.get("sessions"), dl.get("sessions")],
            ["Engaged sessions", a["engagedSessions"], p.get("engagedSessions"), dl.get("engagedSessions")],
            ["Engagement rate", a["engagementRate"], p.get("engagementRate"), None],
            ["Thời gian tương tác TB/phiên (giây)", a["avgEngagementTimePerSession"],
             p.get("avgEngagementTimePerSession"), dl.get("avgEngagementTimePerSession")],
            ["Key events", a["keyEvents"], p.get("keyEvents"), dl.get("keyEvents")],
            ["Session key event rate", a["sessionKeyEventRate"], p.get("sessionKeyEventRate"), None],
            ["Key events / 100 phiên", a["keyEventsPer100Sessions"], p.get("keyEventsPer100Sessions"), None],
            ["Tỷ lệ AI trên tổng phiên site", sh["ofAllSessions"], None, None],
            ["Tổng phiên site", sh["allSessions"], None, None],
            ["Tỷ lệ AI trên nhóm organic", sh["ofOrganicBase"], None, None],
            ["Phiên nhóm organic (AI + Organic Search + Social)", sh["organicBase"], None, None],
            ["Phiên AI trả phí (không tính ở trên)", s["data"]["paidAi"]["sessions"], None, None],
        ]
        sheet = _sheet(wb, "Tổng quan", [("Chỉ số", 48, None), ("Kỳ này", 14, None), ("Kỳ trước", 14, None),
                                         ("% thay đổi", 12, CHG)], rows)
        for i, (label, *_rest) in enumerate(rows, start=2):
            fmt = PCT if ("rate" in label or "Tỷ lệ" in label) else (SEC if "giây" in label or "100" in label else INT)
            sheet.cell(i, 2).number_format = sheet.cell(i, 3).number_format = fmt
        _sheet(wb, "Kênh", [("Kênh", 20, None), ("Phiên", 10, INT), ("Engaged", 10, INT), ("Engagement rate", 14, PCT),
                            ("TG tương tác TB/phiên (giây)", 14, SEC), ("Key events", 10, INT),
                            ("Session key event rate", 14, PCT)],
               [[name, v["sessions"], v["engagedSessions"], v["engagementRate"], v["avgEngagementTimePerSession"],
                 v["keyEvents"], v["sessionKeyEventRate"]]
                for name, v in [("AI Assistants", a)] + list(s["data"]["channels"].items())])

    if src.get("data"):
        _sheet(wb, "Nguồn AI", [("Nguồn", 24, None), ("Phiên", 10, INT), ("Phiên kỳ trước", 12, INT),
                                ("% thay đổi", 12, CHG), ("Engagement rate", 14, PCT),
                                ("TG tương tác TB/phiên (giây)", 14, SEC), ("Key events", 10, INT),
                                ("Session key event rate", 14, PCT), ("Cách nhận diện", 40, None)],
               [[F.SOURCE_NAMES.get(r["source"], r["source"]), r["sessions"], r.get("sessionsPrev"),
                 r.get("sessionsChange"), r.get("engagementRate"), r.get("avgEngagementTimePerSession"), r["keyEvents"],
                 r.get("sessionKeyEventRate"), " + ".join(F.METHOD_NAMES.get(m, m) for m in r["methods"] if m)]
                for r in src["data"]["sources"]])
        _sheet(wb, "Nguồn chi tiết", [("sessionSource", 28, None), ("sessionMedium", 18, None), ("Kênh GA4", 18, None),
                                      ("Phiên", 10, INT), ("Key events", 10, INT), ("Nhận là", 24, None),
                                      ("Cách nhận diện", 32, None)],
               [[r["sessionSource"], r["sessionMedium"], r["sessionDefaultChannelGroup"], r["sessions"], r["keyEvents"],
                 F.SOURCE_NAMES.get(r["subsource"], r["subsource"]), F.METHOD_NAMES.get(r["method"], r["method"])]
                for r in src["data"]["rawRows"]])

    if ke and ke.get("data"):
        _sheet(wb, "Key events", [("Event", 24, None), ("Số lần (phiên AI)", 16, INT),
                                  ("% phiên AI có event", 18, PCT)],
               [[e["event"], e["keyEvents"], e["sessionKeyEventRate"]] for e in ke["data"]["events"]])

    if tr.get("data"):
        key = "week" if tr["data"]["interval"] == "week" else "date"
        series = tr["data"]["series"]
        _sheet(wb, "Xu hướng", [("Tuần (thứ Hai)" if key == "week" else "Ngày", 14, None)] +
               [(F.SOURCE_NAMES.get(x, x), 14, INT) for x in series] +
               ([("Tuần không trọn vẹn", 12, None)] if key == "week" else []),
               [[p[key]] + [p.get(x, 0) for x in series] + (["Có" if p.get("partial") else ""] if key == "week" else [])
                for p in tr["data"]["points"]])

    if land.get("pages", {}).get("data"):
        types, organic = land["types"], land["organic"]
        _sheet(wb, "Loại trang", [("Loại trang", 26, None), ("Số trang", 9, INT), ("Phiên AI", 10, INT),
                                  ("% thay đổi", 12, CHG), ("Engagement rate", 14, PCT), ("Key events", 10, INT),
                                  ("Session key event rate", 14, PCT)],
               [[x["segment"], x["pages"], x["sessions"], x["sessionsChange"], x["engagementRate"], x["keyEvents"],
                 x["sessionKeyEventRate"]] for x in land["segments"]["data"].get("segments", [])])
        _sheet(wb, "Trang đích", [("Trang đích (GA4)", 60, None), ("Loại trang", 22, None), ("Phiên AI", 10, INT),
                                  ("Phiên kỳ trước", 12, INT), ("% thay đổi", 12, CHG), ("Engagement rate", 14, PCT),
                                  ("Key events", 10, INT), ("Session key event rate", 14, PCT),
                                  ("Phiên Organic Search (tham khảo)", 16, INT)],
               [[r["landingPage"], types.get(r["landingPage"], ""), r["sessions"], r["sessionsPrev"],
                 r["sessionsChange"], r["engagementRate"], r["keyEvents"], r["sessionKeyEventRate"],
                 organic.get(r["landingPage"], 0)] for r in land["pages"]["data"]["rows"]])

    if geo.get("countries", {}).get("data"):
        _sheet(wb, "Quốc gia", [("Quốc gia", 22, None), ("Phiên AI", 10, INT), ("Phiên kỳ trước", 12, INT),
                                ("% thay đổi", 12, CHG), ("Engagement rate", 14, PCT), ("Key events", 10, INT),
                                ("Session key event rate", 14, PCT)],
               [[r["country"], r["sessions"], r["sessionsPrev"], r["sessionsChange"], r["engagementRate"],
                 r["keyEvents"], r["sessionKeyEventRate"]] for r in geo["countries"]["data"]["rows"]])
        m = geo["matrix"]["data"].get("matrix", {})
        _sheet(wb, "Quốc gia × nguồn", [("Quốc gia", 22, None), ("Nguồn AI", 24, None), ("Phiên", 10, INT)],
               [[c, F.SOURCE_NAMES.get(sub, sub), n] for c, row in sorted(m.items(), key=lambda kv: -sum(kv[1].values()))
                for sub, n in sorted(row.items(), key=lambda kv: -kv[1])])

    opp = state.get("opp") or {}
    if opp.get("rows") is not None:
        from .ai_opportunities import OPP_FIELDS
        _sheet(wb, "Cơ hội - gợi ý", [(k, 60 if k in ("verification", "action", "url") else 22, None)
                                       for k in OPP_FIELDS],
               [[o[k] for k in OPP_FIELDS] for o in opp.get("opportunities", [])])
        _sheet(wb, "Cơ hội - URL", [("URL", 60, None), ("Loại trang", 22, None), ("Ghép nguồn", 12, None),
                                    ("Đường dẫn GA4", 40, None), ("GA4 AI · Phiên", 12, INT),
                                    ("GA4 AI · Engagement rate", 14, PCT), ("GA4 AI · Key events", 12, INT),
                                    ("GSC Web · Click", 12, INT), ("GSC Web · Impression", 14, INT),
                                    ("GSC Web · CTR", 10, PCT), ("GSC Web · Vị trí TB", 12, "0.0")],
               [[r["url"], r["type"], r["match"], " ".join(r["ga4_paths"]), r["ai_sessions"], r["ai_engagement_rate"],
                 r["ai_key_events"], r["gsc_clicks"], r["gsc_impressions"], r["gsc_ctr"], r["gsc_position"]]
                for r in opp["rows"]])
        ws.append([])
        ws.append(["Cột GSC Web", "Toàn bộ Google Web Search (gồm cả nhưng không tách được AI Overviews/AI Mode): "
                                  "không phải số liệu AI. GSC theo giờ Pacific, GA4 theo giờ Việt Nam."])

    if cfg.db_path.exists():
        _aio_sheets(wb, state, cfg)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _aio_sheets(wb, state: dict, cfg) -> None:
    """Impression AI Overviews / AI Mode từ file GSC đã nhập: theo ngày trong kỳ đang chọn + theo trang của lần nhập
    mới nhất. Ngày không có trong file nào thì không ghi dòng (không ghi 0)."""
    from contextlib import closing

    from . import ai_gsc_import as G
    from . import db
    with closing(db.connect(cfg.db_path)) as con:
        imps = G.imports(con, cfg.site_url)
        if not imps:
            return
        days = G.daily(con, cfg.site_url, state["start"], state["end"])
        web = G.web_daily(con, cfg.site_url, state["start"], state["end"])
        latest = imps[0]
        pages = G.breakdown(con, latest["id"], "page")
    _sheet(wb, "AI Overviews - ngày", [("Ngày (giờ Pacific)", 16, None), ("Impression AI Overviews / AI Mode", 20, INT),
                                        ("Impression GSC Web", 16, INT)],
           [[d, v, web.get(d)] for d, v in sorted(days.items())])
    _sheet(wb, "AI Overviews - trang", [("URL", 70, None), (f"Impression {latest['period_start']}→"
                                                            f"{latest['period_end']}", 20, INT)],
           [[r["key"], r["impressions"]] for r in pages])
